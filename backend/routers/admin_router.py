from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Body
from sqlalchemy.orm import Session
from typing import List, Optional
from datetime import datetime
from io import BytesIO
import openpyxl
import re
import urllib.request
import json
import os
from google.oauth2 import service_account

from database import get_db
import models, schemas, auth
from services.admin_service import AdminService
from services.google_sheets_service import GoogleSheetsService

router = APIRouter(prefix="/admin", tags=["Admin"])



class ImportLinkRequest(schemas.BaseModel if hasattr(schemas, 'BaseModel') else object):
    url: str
    sheet_name: Optional[str] = None


@router.get("/users", response_model=List[schemas.UserResponse])
def list_users(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.list_interns(db)


@router.post("/users", response_model=schemas.UserResponse)
def create_user(
    data: schemas.UserCreate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.create_intern(data, db)


@router.put("/users/{user_id}", response_model=schemas.UserResponse)
def update_user(
    user_id: int,
    data: schemas.UserUpdate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.update_user(user_id, data, db)


@router.delete("/users/{user_id}")
def delete_user(
    user_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.require_admin),
):
    return AdminService.delete_user(user_id, current_user, db)


@router.get("/users/export")
def export_interns(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.export_interns(db)


@router.post("/users/auto-create-sheet")
def auto_create_interns_sheet(
    data: Optional[dict] = Body(default={}),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    target_url = (data or {}).get("target_sheet_url")
    return GoogleSheetsService.auto_create_interns_sheet(db, target_sheet_url=target_url)


@router.get("/users/import-template")
def download_import_template(
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.download_intern_template()


@router.post("/users/import")
def import_users(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Chỉ hỗ trợ định dạng .xlsx")
    try:
        content = file.file.read()
        wb = openpyxl.load_workbook(filename=BytesIO(content), data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Không thể đọc file Excel")
    return AdminService.process_import_users(ws, db)


@router.post("/users/import-link")
def import_users_from_link(
    data: dict,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    url = data.get("url", "")
    match = re.search(r'/d/([a-zA-Z0-9-_]+)', url)
    if not match:
        raise HTTPException(status_code=400, detail="Đường dẫn Google Sheets không hợp lệ")
    sheet_id = match.group(1)
    export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    try:
        req = urllib.request.Request(export_url)
        with urllib.request.urlopen(req) as response:
            content = response.read()
            wb = openpyxl.load_workbook(filename=BytesIO(content), data_only=True)
            ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Không thể tải hoặc đọc dữ liệu từ link.")
    return AdminService.process_import_users(ws, db)


@router.patch("/users/{user_id}/lock")
def toggle_lock(
    user_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.toggle_lock(user_id, db)


@router.patch("/users/{user_id}/reset-password")
def reset_password(
    user_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.reset_password(user_id, db)


@router.get("/accounts", response_model=List[schemas.AdminAccountRow])
def list_accounts(
    user_type: str = Query("intern"),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.list_accounts(user_type, db)


@router.get("/accounts/export")
def export_accounts(
    user_type: str = Query("intern"),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.export_accounts(user_type, db)


@router.get("/stats")
def get_stats(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.get_stats(db)


@router.get("/periods", response_model=List[schemas.PeriodResponse])
def list_periods(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.list_periods(db)


@router.post("/periods", response_model=schemas.PeriodResponse)
def create_period(
    data: schemas.PeriodCreate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.create_period(data, db)


@router.put("/periods/{period_id}", response_model=schemas.PeriodResponse)
def update_period(
    period_id: int,
    data: schemas.PeriodUpdate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.update_period(period_id, data, db)


@router.delete("/periods/{period_id}")
def delete_period(
    period_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.delete_period(period_id, db)


@router.get("/schedule")
def admin_view_schedule(
    month: int = Query(...),
    year: int = Query(...),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.admin_view_schedule(month, year, db)


@router.get("/schedule/export")
def export_schedule(
    month: int = Query(...),
    year: int = Query(...),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return AdminService.export_schedule(month, year, db)


@router.post("/schedule/import")
def import_schedule(
    month: int = Query(...),
    year: int = Query(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Chỉ hỗ trợ file định dạng .xlsx")
    try:
        content = file.file.read()
        wb = openpyxl.load_workbook(filename=BytesIO(content), data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Không thể đọc dữ liệu file Excel")
    return AdminService.process_import_schedule(ws, month, year, db)


@router.post("/schedule/import-link")
def import_schedule_link(
    data: dict,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    from services.google_sheets_service import GoogleSheetsService

    url = data.get("url", "")
    month = data.get("month", datetime.now().month)
    year = data.get("year", datetime.now().year)

    match = re.search(r'/d/([a-zA-Z0-9-_]+)', url)
    if not match:
        raise HTTPException(status_code=400, detail="Đường dẫn Google Sheets không hợp lệ")
    sheet_id = match.group(1)

    # 1. Try reading sheet values via Google API directly using OAuth Credentials
    rows = GoogleSheetsService.read_sheet_values(sheet_id)
    if rows:
        return AdminService.process_import_schedule(rows, month, year, db)

    # 2. Fallback to export URL download if API read unavailable
    export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    try:
        req = urllib.request.Request(export_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            content = response.read()
            wb = openpyxl.load_workbook(filename=BytesIO(content), data_only=True)
            ws = wb.active
            return AdminService.process_import_schedule(ws, month, year, db)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Không thể tải dữ liệu từ Google Sheet. Vui lòng kiểm tra lại link hoặc quyền truy cập!")



@router.post("/schedule/auto-create-sheet")
def auto_create_schedule_sheet(
    data: dict,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    from services.google_sheets_service import GoogleSheetsService
    month = data.get("month", datetime.now().month)
    year = data.get("year", datetime.now().year)
    target_sheet_url = data.get("target_sheet_url")
    force_new = data.get("force_new", False)
    return GoogleSheetsService.auto_create_schedule_sheet(month, year, db, target_sheet_url=target_sheet_url, force_new=force_new)



@router.get("/config/google-json-status")
def get_google_json_status(_: models.User = Depends(auth.require_admin)):
    cfg_path = os.path.join(os.path.dirname(__file__), "..", "config", "google_credentials.json")
    if not os.path.exists(cfg_path):
        return {"configured": False, "client_email": None, "drive_folder_id": None, "drive_sheet_url": None}

    try:
        with open(cfg_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return {
            "configured": True, 
            "client_email": data.get("client_email"),
            "drive_folder_id": data.get("drive_folder_id"),
            "drive_folder_url": data.get("drive_folder_url"),
            "drive_sheet_url": data.get("drive_sheet_url")
        }
    except Exception:
        return {"configured": True, "client_email": None, "drive_folder_id": None, "drive_sheet_url": None}


@router.post("/config/google-json")
@router.post("/config/google-json/")
def save_google_json(
    data: dict,
    _: models.User = Depends(auth.require_admin)
):
    json_text = data.get("json_text", "")
    folder_url_input = data.get("drive_folder_url", "")
    sheet_url_input = data.get("drive_sheet_url", "")

    cfg_dir = os.path.join(os.path.dirname(__file__), "..", "config")
    os.makedirs(cfg_dir, exist_ok=True)
    cfg_path = os.path.join(cfg_dir, "google_credentials.json")

    existing_json = {}
    if os.path.exists(cfg_path):
        try:
            with open(cfg_path, "r", encoding="utf-8") as f:
                existing_json = json.load(f)
        except Exception:
            pass

    if json_text and json_text.strip():
        try:
            json_obj = json.loads(json_text)
        except Exception:
            raise HTTPException(status_code=400, detail="Cú pháp JSON không hợp lệ! Vui lòng kiểm tra lại nội dung dán.")

        required_keys = ["type", "project_id", "private_key", "client_email"]
        missing = [k for k in required_keys if k not in json_obj]
        if missing:
            raise HTTPException(
                status_code=400,
                detail=f"File JSON thiếu các trường bắt buộc của Google Service Account: {', '.join(missing)}"
            )

        if json_obj.get("type") != "service_account":
            raise HTTPException(
                status_code=400,
                detail="File JSON này không phải là chìa khóa 'service_account' của Google Cloud!"
            )

        try:
            service_account.Credentials.from_service_account_info(
                json_obj,
                scopes=['https://www.googleapis.com/auth/spreadsheets', 'https://www.googleapis.com/auth/drive']
            )
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Chìa khóa Google không thể xác thực: {str(e)}")

        existing_json.update(json_obj)

    # Extract folder ID if provided
    if folder_url_input:
        match = re.search(r'/folders/([a-zA-Z0-9-_]+)', folder_url_input)
        if match:
            existing_json['drive_folder_id'] = match.group(1)
            existing_json['drive_folder_url'] = folder_url_input.strip()
        else:
            existing_json['drive_folder_id'] = folder_url_input.strip()
            existing_json['drive_folder_url'] = f"https://drive.google.com/drive/folders/{folder_url_input.strip()}"

    if sheet_url_input:
        existing_json['drive_sheet_url'] = sheet_url_input.strip()

    if not existing_json:
        raise HTTPException(status_code=400, detail="Nội dung JSON, Link Thư mục hoặc Link Google Sheet không được để trống!")

    with open(cfg_path, "w", encoding="utf-8") as f:
        json.dump(existing_json, f, indent=2)

    client_email = existing_json.get("client_email", "")
    return {
        "success": True,
        "valid": True,
        "client_email": client_email,
        "drive_folder_id": existing_json.get("drive_folder_id"),
        "drive_sheet_url": existing_json.get("drive_sheet_url"),
        "message": f"✓ Đã lưu cấu hình Google Service Account & Google Sheet/Drive thành công!"
    }


@router.post("/config/env-json")
@router.post("/config/env-json/")
def save_env_json(
    data: dict,
    _: models.User = Depends(auth.require_admin)
):
    json_text = data.get("json_text", "")
    if not json_text or not json_text.strip():
        raise HTTPException(status_code=400, detail="Vui lòng dán hoặc chọn file JSON!")

    try:
        parsed = json.loads(json_text)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Cú pháp JSON không hợp lệ: {str(e)}")

    env_vars = {}

    # 1. Google OAuth Web Client JSON format (e.g. from Google Cloud console Download JSON)
    if "web" in parsed or "installed" in parsed:
        client_info = parsed.get("web") or parsed.get("installed") or {}
        env_vars["GOOGLE_CLIENT_ID"] = client_info.get("client_id", "")
        env_vars["GOOGLE_CLIENT_SECRET"] = client_info.get("client_secret", "")
        env_vars["GOOGLE_PROJECT_ID"] = client_info.get("project_id", "")

        # Also sync to google_oauth.json
        oauth_cfg = os.path.join(os.path.dirname(__file__), "..", "config", "google_oauth.json")
        os.makedirs(os.path.dirname(oauth_cfg), exist_ok=True)
        with open(oauth_cfg, "w", encoding="utf-8") as f:
            json.dump({
                "client_id": client_info.get("client_id", ""),
                "client_secret": client_info.get("client_secret", "")
            }, f, indent=2)

    # 2. Google Service Account JSON format
    elif parsed.get("type") == "service_account":
        env_vars["GOOGLE_TYPE"] = parsed.get("type", "")
        env_vars["GOOGLE_PROJECT_ID"] = parsed.get("project_id", "")
        env_vars["GOOGLE_PRIVATE_KEY_ID"] = parsed.get("private_key_id", "")
        env_vars["GOOGLE_PRIVATE_KEY"] = parsed.get("private_key", "").replace("\n", "\\n")
        env_vars["GOOGLE_CLIENT_EMAIL"] = parsed.get("client_email", "")

        # Also sync to google_credentials.json
        cred_cfg = os.path.join(os.path.dirname(__file__), "..", "config", "google_credentials.json")
        os.makedirs(os.path.dirname(cred_cfg), exist_ok=True)
        with open(cred_cfg, "w", encoding="utf-8") as f:
            json.dump(parsed, f, indent=2)

    # 3. Generic Key-Value JSON
    else:
        for k, v in parsed.items():
            if isinstance(v, (str, int, float, bool)):
                env_vars[k.upper()] = str(v)

    if not env_vars:
        raise HTTPException(status_code=400, detail="Không tìm thấy dữ liệu hợp lệ trong file JSON!")

    # Write to root .env and backend .env
    root_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
    backend_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))

    target_env_paths = [
        os.path.join(root_dir, ".env"),
        os.path.join(backend_dir, ".env")
    ]

    for env_path in target_env_paths:
        existing_env = {}
        if os.path.exists(env_path):
            with open(env_path, "r", encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if line and not line.startswith("#") and "=" in line:
                        k, v = line.split("=", 1)
                        existing_env[k.strip()] = v.strip()

        existing_env.update(env_vars)

        with open(env_path, "w", encoding="utf-8") as f:
            f.write("# Google OAuth & Environment Config\n")
            for k, v in existing_env.items():
                f.write(f"{k}={v}\n")

    return {
        "success": True,
        "env_vars": env_vars,
        "message": f"✓ Đã phân tích JSON và tự động lưu {len(env_vars)} biến môi trường vào file .env thành công!"
    }






