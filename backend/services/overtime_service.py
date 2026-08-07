from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import List, Optional
from datetime import date, datetime
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import calendar

import models
import schemas
import auth

SPLIT_HOUR = 22 * 60  # 22:00 in minutes


def parse_time_minutes(t: str) -> int:
    h, m = map(int, t.split(":"))
    return h * 60 + m


def minutes_to_hhmm(mins: int) -> str:
    return f"{mins // 60:02d}:{mins % 60:02d}"


def calc_raw_hours(start_min: int, end_min: int) -> float:
    return round((end_min - start_min) / 60, 2)


def get_factor(is_holiday: bool, is_weekend: bool, before_22: bool) -> float:
    if is_holiday:
        return 3.0 if before_22 else 3.9
    if is_weekend:
        return 2.0 if before_22 else 2.7
    return 1.5 if before_22 else 2.1


def split_segments(start_min: int, end_min: int, is_holiday: bool, is_weekend: bool) -> list:
    segments = []
    if start_min < SPLIT_HOUR and end_min > SPLIT_HOUR:
        seg1_raw = calc_raw_hours(start_min, SPLIT_HOUR)
        seg1_factor = get_factor(is_holiday, is_weekend, before_22=True)
        segments.append({
            "start_time": minutes_to_hhmm(start_min),
            "end_time": minutes_to_hhmm(SPLIT_HOUR),
            "raw_hours": seg1_raw,
            "factor": seg1_factor,
            "weighted_hours": round(seg1_raw * seg1_factor, 2),
        })
        seg2_raw = calc_raw_hours(SPLIT_HOUR, end_min)
        seg2_factor = get_factor(is_holiday, is_weekend, before_22=False)
        segments.append({
            "start_time": minutes_to_hhmm(SPLIT_HOUR),
            "end_time": minutes_to_hhmm(end_min),
            "raw_hours": seg2_raw,
            "factor": seg2_factor,
            "weighted_hours": round(seg2_raw * seg2_factor, 2),
        })
    else:
        before_22 = end_min <= SPLIT_HOUR
        raw = calc_raw_hours(start_min, end_min)
        factor = get_factor(is_holiday, is_weekend, before_22=before_22)
        segments.append({
            "start_time": minutes_to_hhmm(start_min),
            "end_time": minutes_to_hhmm(end_min),
            "raw_hours": raw,
            "factor": factor,
            "weighted_hours": round(raw * factor, 2),
        })
    return segments


def validate_ot_time(start: str, end: str, work_date: date):
    try:
        start_min = parse_time_minutes(start)
        end_min = parse_time_minutes(end)
    except Exception:
        raise HTTPException(status_code=400, detail="Định dạng giờ không hợp lệ. Dùng HH:MM")
    if work_date.weekday() in (0, 1, 2, 3, 4):
        if start_min < 18 * 60 + 30:
            raise HTTPException(
                status_code=400,
                detail="Thời gian bắt đầu OT các ngày từ Thứ 2 đến Thứ 6 phải từ 18:30 trở đi."
            )
    if end_min <= start_min:
        raise HTTPException(
            status_code=400,
            detail="Thời gian kết thúc phải lớn hơn thời gian bắt đầu."
        )
    raw = calc_raw_hours(start_min, end_min)
    if raw > 24:
        raise HTTPException(status_code=400, detail="OT không được quá 24 giờ.")
    return start_min, end_min, raw


def check_overlap(
    db: Session,
    user_id: int,
    work_date: date,
    start_min: int,
    end_min: int,
    exclude_ids: Optional[list] = None,
):
    existing_q = db.query(models.OvertimeRequest).filter(
        models.OvertimeRequest.user_id == user_id,
        models.OvertimeRequest.work_date == work_date,
        models.OvertimeRequest.status != "Rejected",
    )
    if exclude_ids:
        existing_q = existing_q.filter(
            models.OvertimeRequest.id.notin_(exclude_ids)
        )
    for ot in existing_q.all():
        ex_start = parse_time_minutes(ot.start_time)
        ex_end = parse_time_minutes(ot.end_time)
        if start_min < ex_end and end_min > ex_start:
            raise HTTPException(
                status_code=400,
                detail=f"Khoảng thời gian OT bị trùng với đăng ký đã có ({ot.start_time}–{ot.end_time})."
            )


def enrich_ot(ot: models.OvertimeRequest) -> dict:
    d = {c.name: getattr(ot, c.name) for c in ot.__table__.columns}
    if ot.user:
        d["employee_code"] = ot.user.employee_code
        d["full_name"] = ot.user.full_name
    return d


class OvertimeService:
    @staticmethod
    def create_overtime(body: schemas.OvertimeCreate, current_user: models.User, db: Session) -> dict:
        if current_user.user_type != "employee":
            raise HTTPException(status_code=403, detail="Chỉ nhân sự mới được đăng ký OT.")
        if current_user.staff_category != "Onsite":
            raise HTTPException(status_code=403, detail="Bạn chưa được cấp quyền đăng ký OT. Vui lòng liên hệ quản trị viên.")
        if not current_user.project:
            raise HTTPException(status_code=403, detail="Bạn chưa được gán dự án. Vui lòng liên hệ quản trị viên.")
        if current_user.account_status != 1:
            raise HTTPException(status_code=403, detail="Tài khoản đang bị khóa.")

        start_min, end_min, _ = validate_ot_time(body.start_time, body.end_time, body.work_date)
        check_overlap(db, current_user.id, body.work_date, start_min, end_min)

        is_weekend = body.work_date.weekday() in (5, 6)
        segments = split_segments(start_min, end_min, body.is_holiday, is_weekend)

        created = []
        for seg in segments:
            ot = models.OvertimeRequest(
                user_id=current_user.id,
                project=current_user.project,
                work_date=body.work_date,
                start_time=seg["start_time"],
                end_time=seg["end_time"],
                raw_hours=seg["raw_hours"],
                factor=seg["factor"],
                weighted_hours=seg["weighted_hours"],
                reason=body.reason,
                status="Pending",
            )
            db.add(ot)
            db.flush()
            created.append(ot)

        db.commit()
        for ot in created:
            db.refresh(ot)

        return {
            "segments": len(segments),
            "records": [enrich_ot(ot) for ot in created],
            "total_raw_hours": round(sum(s["raw_hours"] for s in segments), 2),
            "total_weighted_hours": round(sum(s["weighted_hours"] for s in segments), 2),
        }

    @staticmethod
    def get_my_overtime(month: Optional[int], year: Optional[int], current_user: models.User, db: Session) -> List[dict]:
        now = datetime.now()
        m = month or now.month
        y = year or now.year
        first_day = date(y, m, 1)
        last_day = date(y, m, calendar.monthrange(y, m)[1])

        records = db.query(models.OvertimeRequest).filter(
            models.OvertimeRequest.user_id == current_user.id,
            models.OvertimeRequest.work_date >= first_day,
            models.OvertimeRequest.work_date <= last_day,
        ).order_by(models.OvertimeRequest.work_date).all()

        return [enrich_ot(r) for r in records]

    @staticmethod
    def get_my_ot_stats(month: Optional[int], year: Optional[int], current_user: models.User, db: Session) -> dict:
        now = datetime.now()
        m = month or now.month
        y = year or now.year

        first_day = date(y, m, 1)
        last_day = date(y, m, calendar.monthrange(y, m)[1])

        records = db.query(models.OvertimeRequest).filter(
            models.OvertimeRequest.user_id == current_user.id,
            models.OvertimeRequest.work_date >= first_day,
            models.OvertimeRequest.work_date <= last_day,
        ).all()

        total_raw = sum(r.raw_hours for r in records if r.status == "Approved")
        total_weighted = sum(r.weighted_hours for r in records if r.status == "Approved")
        pending = sum(1 for r in records if r.status == "Pending")
        approved = sum(1 for r in records if r.status == "Approved")
        rejected = sum(1 for r in records if r.status == "Rejected")

        return {
            "month": m,
            "year": y,
            "total_raw_hours": round(total_raw, 2),
            "total_weighted_hours": round(total_weighted, 2),
            "pending_count": pending,
            "approved_count": approved,
            "rejected_count": rejected,
            "is_onsite": current_user.staff_category == "Onsite",
            "has_project": bool(current_user.project),
        }

    @staticmethod
    def preview_ot_split(body: schemas.OvertimeCreate) -> dict:
        try:
            start_min = parse_time_minutes(body.start_time)
            end_min = parse_time_minutes(body.end_time)
        except Exception:
            raise HTTPException(status_code=400, detail="Định dạng giờ không hợp lệ.")
        if end_min <= start_min:
            raise HTTPException(status_code=400, detail="Giờ kết thúc phải lớn hơn giờ bắt đầu.")

        is_weekend = body.work_date.weekday() in (5, 6)
        segments = split_segments(start_min, end_min, body.is_holiday, is_weekend)
        return {
            "segments": segments,
            "total_raw_hours": round(sum(s["raw_hours"] for s in segments), 2),
            "total_weighted_hours": round(sum(s["weighted_hours"] for s in segments), 2),
            "is_weekend": is_weekend,
            "is_holiday": body.is_holiday,
        }

    @staticmethod
    def update_overtime(ot_id: int, body: schemas.OvertimeUpdate, current_user: models.User, db: Session) -> dict:
        ot = db.query(models.OvertimeRequest).filter(
            models.OvertimeRequest.id == ot_id,
            models.OvertimeRequest.user_id == current_user.id,
        ).first()
        if not ot:
            raise HTTPException(status_code=404, detail="Không tìm thấy yêu cầu OT.")
        if ot.status == "Approved":
            raise HTTPException(status_code=400, detail="Không thể sửa OT đã được duyệt.")

        new_start = body.start_time or ot.start_time
        new_end = body.end_time or ot.end_time
        is_holiday = body.is_holiday if body.is_holiday is not None else False

        start_min, end_min, _ = validate_ot_time(new_start, new_end, ot.work_date)
        check_overlap(db, current_user.id, ot.work_date, start_min, end_min, exclude_ids=[ot_id])

        is_weekend = ot.work_date.weekday() in (5, 6)
        before_22 = end_min <= SPLIT_HOUR
        new_factor = get_factor(is_holiday, is_weekend, before_22=before_22)
        raw_hours = calc_raw_hours(start_min, end_min)

        ot.start_time = new_start
        ot.end_time = new_end
        ot.factor = new_factor
        ot.raw_hours = raw_hours
        ot.weighted_hours = round(raw_hours * new_factor, 2)
        if body.reason is not None:
            ot.reason = body.reason
        ot.status = "Pending"
        ot.reject_reason = None
        db.commit()
        db.refresh(ot)
        return enrich_ot(ot)

    @staticmethod
    def delete_overtime(ot_id: int, current_user: models.User, db: Session) -> dict:
        ot = db.query(models.OvertimeRequest).filter(
            models.OvertimeRequest.id == ot_id,
            models.OvertimeRequest.user_id == current_user.id,
        ).first()
        if not ot:
            raise HTTPException(status_code=404, detail="Không tìm thấy yêu cầu OT.")
        if ot.status == "Approved":
            raise HTTPException(status_code=400, detail="Không thể xóa OT đã được duyệt.")
        db.delete(ot)
        db.commit()
        return {"message": "Đã xóa yêu cầu OT."}

    @staticmethod
    def admin_list_ot(
        month: Optional[int],
        year: Optional[int],
        project: Optional[str],
        status: Optional[str],
        employee_name: Optional[str],
        staff_category: Optional[str],
        db: Session
    ) -> List[dict]:
        now = datetime.now()
        m = month or now.month
        y = year or now.year
        first_day = date(y, m, 1)
        last_day = date(y, m, calendar.monthrange(y, m)[1])

        q = db.query(models.OvertimeRequest).join(
            models.User, models.OvertimeRequest.user_id == models.User.id
        ).filter(
            models.OvertimeRequest.work_date >= first_day,
            models.OvertimeRequest.work_date <= last_day,
        )

        if project:
            q = q.filter(models.OvertimeRequest.project.ilike(f"%{project}%"))
        if status:
            q = q.filter(models.OvertimeRequest.status == status)
        if employee_name:
            q = q.filter(models.User.full_name.ilike(f"%{employee_name}%"))
        if staff_category:
            q = q.filter(models.User.staff_category == staff_category)

        records = q.order_by(models.OvertimeRequest.work_date, models.User.full_name).all()
        return [enrich_ot(r) for r in records]

    @staticmethod
    def admin_ot_summary(month: Optional[int], year: Optional[int], project: Optional[str], db: Session) -> dict:
        now = datetime.now()
        m = month or now.month
        y = year or now.year
        first_day = date(y, m, 1)
        last_day = date(y, m, calendar.monthrange(y, m)[1])
        num_days = calendar.monthrange(y, m)[1]

        # Query all OT requests in the month (or for specific project)
        q = db.query(models.OvertimeRequest).join(
            models.User, models.OvertimeRequest.user_id == models.User.id
        ).filter(
            models.OvertimeRequest.work_date >= first_day,
            models.OvertimeRequest.work_date <= last_day,
        )
        if project:
            q = q.filter(models.OvertimeRequest.project.ilike(f"%{project}%"))

        records = q.order_by(models.User.full_name, models.OvertimeRequest.work_date, models.OvertimeRequest.start_time).all()
        factors_order = [1.5, 2.1, 2.0, 2.7, 3.0, 3.9]
        user_map: dict = {}
        user_order: list = []

        day_of_week_names = {
            0: "Thứ 2",
            1: "Thứ 3",
            2: "Thứ 4",
            3: "Thứ 5",
            4: "Thứ 6",
            5: "Thứ 7",
            6: "Chủ nhật"
        }

        for ot in records:
            uid = ot.user_id
            if uid not in user_map:
                user_map[uid] = {
                    "user_id": uid,
                    "employee_code": ot.user.employee_code if ot.user else "",
                    "full_name": ot.user.full_name if ot.user else "",
                    "project": ot.project or (ot.user.project if ot.user else ""),
                    "days": {},
                    "detail_records": [],
                    "total_by_factor": {f"{f:.1f}": 0.0 for f in factors_order},
                    "total_raw": 0.0,
                    "total_weighted": 0.0,
                    "total_raw_hours": 0.0,
                    "total_weighted_hours": 0.0,
                    "approved_raw_hours": 0.0,
                    "approved_weighted_hours": 0.0,
                    "pending_raw_hours": 0.0,
                    "pending_weighted_hours": 0.0,
                    "rejected_raw_hours": 0.0,
                    "rejected_weighted_hours": 0.0,
                    "total_count": 0,
                    "pending_count": 0,
                    "approved_count": 0,
                    "rejected_count": 0,
                }
                user_order.append(uid)

            dow_idx = ot.work_date.weekday()
            dow_name = day_of_week_names.get(dow_idx, f"Thứ {dow_idx + 2}")

            record_item = {
                "ot_id": ot.id,
                "work_date": ot.work_date.strftime("%Y-%m-%d"),
                "date_display": ot.work_date.strftime("%d/%m/%Y"),
                "day_number": ot.work_date.day,
                "day_of_week": dow_name,
                "start_time": ot.start_time,
                "end_time": ot.end_time,
                "raw_hours": round(float(ot.raw_hours or 0.0), 2),
                "factor": float(ot.factor or 1.5),
                "weighted_hours": round(float(ot.weighted_hours or 0.0), 2),
                "reason": ot.reason or "",
                "status": ot.status,
                "reject_reason": ot.reject_reason or "",
                "is_holiday": float(ot.factor or 0) >= 3.0,
                "is_weekend": dow_idx in (5, 6),
            }

            day = ot.work_date.day
            if day not in user_map[uid]["days"]:
                user_map[uid]["days"][day] = []
            user_map[uid]["days"][day].append(record_item)
            user_map[uid]["detail_records"].append(record_item)

            factor_key = f"{float(ot.factor):.1f}"
            raw = round(float(ot.raw_hours or 0.0), 2)
            weighted = round(float(ot.weighted_hours or 0.0), 2)

            user_map[uid]["total_count"] += 1
            user_map[uid]["total_raw"] = round(user_map[uid]["total_raw"] + raw, 2)
            user_map[uid]["total_weighted"] = round(user_map[uid]["total_weighted"] + weighted, 2)
            user_map[uid]["total_raw_hours"] = user_map[uid]["total_raw"]
            user_map[uid]["total_weighted_hours"] = user_map[uid]["total_weighted"]

            user_map[uid]["total_by_factor"][factor_key] = round(
                user_map[uid]["total_by_factor"].get(factor_key, 0.0) + raw, 2
            )

            if ot.status == "Approved":
                user_map[uid]["approved_count"] += 1
                user_map[uid]["approved_raw_hours"] = round(user_map[uid]["approved_raw_hours"] + raw, 2)
                user_map[uid]["approved_weighted_hours"] = round(user_map[uid]["approved_weighted_hours"] + weighted, 2)
            elif ot.status == "Pending":
                user_map[uid]["pending_count"] += 1
                user_map[uid]["pending_raw_hours"] = round(user_map[uid]["pending_raw_hours"] + raw, 2)
                user_map[uid]["pending_weighted_hours"] = round(user_map[uid]["pending_weighted_hours"] + weighted, 2)
            elif ot.status == "Rejected":
                user_map[uid]["rejected_count"] += 1
                user_map[uid]["rejected_raw_hours"] = round(user_map[uid]["rejected_raw_hours"] + raw, 2)
                user_map[uid]["rejected_weighted_hours"] = round(user_map[uid]["rejected_weighted_hours"] + weighted, 2)

        return {
            "month": m,
            "year": y,
            "num_days": num_days,
            "rows": [user_map[uid] for uid in user_order],
        }

    @staticmethod
    def admin_approve_ot(ot_id: int, admin: models.User, db: Session) -> dict:
        ot = db.query(models.OvertimeRequest).filter(models.OvertimeRequest.id == ot_id).first()
        if not ot:
            raise HTTPException(status_code=404, detail="Không tìm thấy yêu cầu OT.")
        if ot.status != "Pending":
            raise HTTPException(status_code=400, detail="Chỉ duyệt được OT đang ở trạng thái Pending.")
        ot.status = "Approved"
        ot.approved_by = admin.id
        ot.approved_at = datetime.now()
        ot.reject_reason = None
        db.commit()
        db.refresh(ot)
        return enrich_ot(ot)

    @staticmethod
    def admin_reject_ot(ot_id: int, body: schemas.OvertimeApprove, admin: models.User, db: Session) -> dict:
        if not body.reject_reason or not body.reject_reason.strip():
            raise HTTPException(status_code=400, detail="Vui lòng nhập lý do từ chối.")
        ot = db.query(models.OvertimeRequest).filter(models.OvertimeRequest.id == ot_id).first()
        if not ot:
            raise HTTPException(status_code=404, detail="Không tìm thấy yêu cầu OT.")
        if ot.status != "Pending":
            raise HTTPException(status_code=400, detail="Chỉ từ chối được OT đang ở trạng thái Pending.")
        ot.status = "Rejected"
        ot.reject_reason = body.reject_reason.strip()
        ot.approved_by = admin.id
        ot.approved_at = datetime.now()
        db.commit()
        db.refresh(ot)
        return enrich_ot(ot)

    @staticmethod
    def admin_approve_selected(body: schemas.OvertimeApproveSelected, admin: models.User, db: Session) -> dict:
        updated = 0
        for ot_id in body.ids:
            ot = db.query(models.OvertimeRequest).filter(
                models.OvertimeRequest.id == ot_id,
                models.OvertimeRequest.status == "Pending",
            ).first()
            if ot:
                ot.status = "Approved"
                ot.approved_by = admin.id
                ot.approved_at = datetime.now()
                updated += 1
        db.commit()
        return {"message": f"Đã duyệt {updated} yêu cầu OT."}

    @staticmethod
    def admin_approve_all_pending(month: Optional[int], year: Optional[int], project: Optional[str], admin: models.User, db: Session) -> dict:
        now = datetime.now()
        m = month or now.month
        y = year or now.year
        first_day = date(y, m, 1)
        last_day = date(y, m, calendar.monthrange(y, m)[1])

        q = db.query(models.OvertimeRequest).filter(
            models.OvertimeRequest.status == "Pending",
            models.OvertimeRequest.work_date >= first_day,
            models.OvertimeRequest.work_date <= last_day,
        )
        if project:
            q = q.filter(models.OvertimeRequest.project.ilike(f"%{project}%"))

        records = q.all()
        for ot in records:
            ot.status = "Approved"
            ot.approved_by = admin.id
            ot.approved_at = datetime.now()
        db.commit()
        return {"message": f"Đã duyệt {len(records)} yêu cầu OT."}

    @staticmethod
    def admin_export_excel(month: int, year: int, project: Optional[str], db: Session) -> StreamingResponse:
        first_day = date(year, month, 1)
        last_day = date(year, month, calendar.monthrange(year, month)[1])
        num_days = calendar.monthrange(year, month)[1]

        q = db.query(models.OvertimeRequest).join(
            models.User, models.OvertimeRequest.user_id == models.User.id
        ).filter(
            models.OvertimeRequest.work_date >= first_day,
            models.OvertimeRequest.work_date <= last_day,
            models.OvertimeRequest.status == "Approved",
        )
        if project:
            q = q.filter(models.OvertimeRequest.project.ilike(f"%{project}%"))

        records = q.order_by(models.User.full_name, models.OvertimeRequest.work_date).all()
        factors_order = [1.5, 2.1, 2.0, 2.7, 3.0, 3.9]

        user_map: dict = {}
        user_order: list = []
        for ot in records:
            uid = ot.user_id
            if uid not in user_map:
                user_map[uid] = {
                    "no": len(user_order) + 1,
                    "employee_code": ot.user.employee_code,
                    "full_name": ot.user.full_name,
                    "project": ot.project or "",
                    "position": ot.user.position or "",
                    "days": {},
                    "by_factor": {round(float(f), 1): 0.0 for f in factors_order},
                    "total_raw": 0.0,
                    "total_weighted": 0.0,
                }
                user_order.append(uid)
            d = ot.work_date.day
            user_map[uid]["days"][d] = round(user_map[uid]["days"].get(d, 0.0) + ot.raw_hours, 2)
            f_norm = round(float(ot.factor), 1)
            user_map[uid]["by_factor"][f_norm] = round(
                user_map[uid]["by_factor"].get(f_norm, 0.0) + ot.raw_hours, 2
            )
            user_map[uid]["total_raw"] = round(user_map[uid]["total_raw"] + ot.raw_hours, 2)
            user_map[uid]["total_weighted"] = round(user_map[uid]["total_weighted"] + ot.weighted_hours, 2)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"OT T{month:02d}.{year}"

        project_name = project or (records[0].project if records else "")
        month_str = f"Tháng {month:02d} Năm {year}"

        GREEN_LIGHT  = "C6EFCE"
        GREEN_HEADER = "92D050"
        YELLOW_FACTOR= "FFEB9C"
        PINK_TOTAL   = "FFC7CE"
        YELLOW_DATA  = "FFFF99"

        ca = Alignment(horizontal="center", vertical="center", wrap_text=True)
        la = Alignment(horizontal="left",   vertical="center", wrap_text=True)
        bold11 = Font(name="Times New Roman", bold=True,  size=11)
        bold13 = Font(name="Times New Roman", bold=True,  size=13)
        bold10 = Font(name="Times New Roman", bold=True,  size=10)
        norm10 = Font(name="Times New Roman", bold=False, size=10)
        sm9    = Font(name="Times New Roman", bold=False, size=9)

        FIXED   = 5
        FAC     = 6
        DAY_S   = FIXED + 1
        FAC_S   = DAY_S + num_days
        TOT_COL = FAC_S + FAC
        TOT_W   = TOT_COL + 1
        LAST    = TOT_W

        for ci, w in [(1,5),(2,11),(3,20),(4,14),(5,12)]:
            ws.column_dimensions[get_column_letter(ci)].width = w
        for ci in range(DAY_S, FAC_S):
            ws.column_dimensions[get_column_letter(ci)].width = 4.5
        for ci in range(FAC_S, LAST + 1):
            ws.column_dimensions[get_column_letter(ci)].width = 8

        def _thin():
            s = Side(border_style="thin", color="000000")
            return Border(left=s, right=s, top=s, bottom=s)

        def mc(r1, c1, r2, c2, val="", font=None, fill=None, align=None, border=None):
            if r1 != r2 or c1 != c2:
                ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
            cell = ws.cell(r1, c1, val)
            if font:   cell.font   = font
            if fill:   cell.fill   = fill
            if align:  cell.alignment = align
            if border: cell.border = border
            return cell

        ws.row_dimensions[1].height = 20
        mc(1,1,1,FIXED+10, "CÔNG TY ĐẦU TƯ CÔNG NGHỆ VIETTEL", bold11, align=la)
        mc(1,FAC_S,1,LAST, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", Font(name="Times New Roman",bold=True,size=11), align=ca)

        ws.row_dimensions[2].height = 18
        mc(2,1,2,FIXED+10, "PHÒNG CHÍNH TRỊ NHÂN SỰ", Font(name="Times New Roman",bold=True,underline="single",size=11), align=la)
        mc(2,FAC_S,2,LAST, "Độc lập - Tự do - Hạnh phúc", Font(name="Times New Roman",bold=False,italic=True,underline="single",size=11), align=ca)

        ws.row_dimensions[3].height = 10
        ws.row_dimensions[4].height = 24
        mc(4,1,4,LAST, "PHỤ LỤC 02: BẢNG TỔNG HỢP CÔNG THỜI GIAN LÀM THÊM GIỜ", bold13, align=ca)

        ws.row_dimensions[5].height = 20
        mc(5,1,5,LAST, f"TÊN DỰ ÁN: {project_name.upper()}", Font(name="Times New Roman",bold=True,size=12), PatternFill(fill_type="solid", fgColor=GREEN_LIGHT), ca)

        ws.row_dimensions[6].height = 18
        mc(6,1,6,LAST, month_str, Font(name="Times New Roman",bold=True,size=12), align=ca)

        ws.row_dimensions[7].height = 32
        ws.row_dimensions[8].height = 42

        fixed_headers = ["STT", "MNV", "Họ và tên", "Phòng ban/Dự án", "Vị trí"]
        for i, h in enumerate(fixed_headers, 1):
            mc(7,i,8,i, h, bold10, PatternFill(fill_type="solid", fgColor=GREEN_HEADER), ca, _thin())

        mc(7, DAY_S, 7, FAC_S-1, f"Ngày trong tháng {month:02d}/{year}", bold10, PatternFill(fill_type="solid", fgColor=GREEN_HEADER), ca, _thin())
        for d in range(1, num_days+1):
            c = ws.cell(8, DAY_S+d-1, d)
            c.font = sm9; c.fill = PatternFill(fill_type="solid", fgColor=GREEN_HEADER)
            c.alignment = ca; c.border = _thin()

        mc(7, FAC_S, 7, FAC_S+FAC-1, "Tổng thêm giờ", bold10, PatternFill(fill_type="solid", fgColor=YELLOW_FACTOR), ca, _thin())
        factor_labels = ["1.5x\n≤22h","2.1x\n>22h","2.0x\n≤22h\nT7,CN","2.7x\n>22h\nT7,CN","3.0x\n≤22h\nLễ","3.9x\n>22h\nLễ"]
        for i, lbl in enumerate(factor_labels):
            c = ws.cell(8, FAC_S+i, lbl)
            c.font = sm9; c.fill = PatternFill(fill_type="solid", fgColor=YELLOW_FACTOR)
            c.alignment = ca; c.border = _thin()

        mc(7, TOT_COL, 8, TOT_COL, "Tổng giờ\nOT (giờ)", bold10, PatternFill(fill_type="solid", fgColor=PINK_TOTAL), ca, _thin())
        mc(7, TOT_W, 8, TOT_W, "Tổng\nquy đổi\n(giờ)", bold10, PatternFill(fill_type="solid", fgColor=PINK_TOTAL), ca, _thin())

        DATA_R = 9
        col_day_totals = {d: 0.0 for d in range(1, num_days+1)}
        col_fac_totals = {f: 0.0 for f in factors_order}
        grand_raw = 0.0
        grand_w   = 0.0

        for ri, uid in enumerate(user_order):
            u = user_map[uid]
            r = DATA_R + ri
            ws.row_dimensions[r].height = 18

            for ci, val in enumerate([u["no"],u["employee_code"],u["full_name"],u["project"],u["position"]], 1):
                c = ws.cell(r, ci, val)
                c.font = norm10; c.fill = PatternFill(fill_type="solid", fgColor=YELLOW_DATA)
                c.alignment = la if ci == 3 else ca
                c.border = _thin()

            for d in range(1, num_days+1):
                col = DAY_S + d - 1
                val = u["days"].get(d, "")
                c = ws.cell(r, col, val if val else "")
                c.font = norm10; c.fill = PatternFill(fill_type="solid", fgColor=YELLOW_DATA)
                c.alignment = ca; c.border = _thin()
                if val:
                    col_day_totals[d] = round(col_day_totals.get(d, 0.0) + val, 2)

            for i, fv in enumerate(factors_order):
                col = FAC_S + i
                val = u["by_factor"].get(fv, 0.0)
                c = ws.cell(r, col, round(val, 2) if val else "")
                c.font = norm10; c.fill = PatternFill(fill_type="solid", fgColor=YELLOW_DATA)
                c.alignment = ca; c.border = _thin()
                col_fac_totals[fv] = round(col_fac_totals.get(fv, 0.0) + val, 2)

            grand_raw += u["total_raw"]
            grand_w   += u["total_weighted"]
            for col, val in [(TOT_COL, u["total_raw"]), (TOT_W, u["total_weighted"])]:
                c = ws.cell(r, col, round(val, 2) if val else "")
                c.font = norm10; c.fill = PatternFill(fill_type="solid", fgColor=YELLOW_DATA)
                c.alignment = ca; c.border = _thin()

        SR = DATA_R + len(user_order)
        ws.row_dimensions[SR].height = 20
        mc(SR,1,SR,FIXED, "Cộng", bold10, PatternFill(fill_type="solid", fgColor=PINK_TOTAL), ca, _thin())
        for d in range(1, num_days+1):
            val = col_day_totals.get(d, 0.0)
            c = ws.cell(SR, DAY_S+d-1, round(val, 2) if val else "")
            c.font = bold10; c.fill = PatternFill(fill_type="solid", fgColor=PINK_TOTAL)
            c.alignment = ca; c.border = _thin()
        for i, fv in enumerate(factors_order):
            val = col_fac_totals.get(fv, 0.0)
            c = ws.cell(SR, FAC_S+i, round(val, 2) if val else "")
            c.font = bold10; c.fill = PatternFill(fill_type="solid", fgColor=PINK_TOTAL)
            c.alignment = ca; c.border = _thin()
        for col, val in [(TOT_COL, grand_raw), (TOT_W, grand_w)]:
            c = ws.cell(SR, col, round(val, 2))
            c.font = bold10; c.fill = PatternFill(fill_type="solid", fgColor=PINK_TOTAL)
            c.alignment = ca; c.border = _thin()

        ws.freeze_panes = ws.cell(DATA_R, DAY_S)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = f"OT_T{month:02d}_{year}.xlsx"
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={fname}"}
        )
