import os
import json
import time
import datetime
import logging
from typing import Dict, Any, List, Optional
from sqlalchemy import text
from fastapi import HTTPException
from google import genai

from schemas import ChatRequest, ChatResponse, RunSQLRequest
from database_utils import db_execute, db_commit
from services.hrai.config import get_settings
from services.hrai.query_planner_engine.engine_pipeline import run_query_planner_pipeline
from services.hrai.ai_pipeline.execution_engine import execute_query, validate_sql
from services.hrai.sheet_executor import create_table_and_sync_sheet
from services.hrai.sheet_pipeline import (
    preprocess_sheet_raw_data,
    analyze_columns_with_gemini_stage1,
    analyze_sections_with_gemini_stage2,
    calculate_dynamic_section_ranges
)

logger = logging.getLogger(__name__)

ENV_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".env"))

SYSTEM_TABLE_NAMES = {
    'users', 'accounts', 'schedules', 'positions', 
    'overtime_requests', 'schedule_period', 'alembic_version'
}


class HraiService:
    @staticmethod
    async def chat(request: ChatRequest, db: Any) -> ChatResponse:
        start_time = time.time()
        try:
            response = await run_query_planner_pipeline(
                question=request.question,
                session_id=request.session_id,
                db=db,
            )
            elapsed = round(time.time() - start_time, 3)
            response.metadata["execution_time_seconds"] = elapsed
            logger.info("Chat completed | elapsed=%.3fs", elapsed)
            return response
        except Exception as e:
            elapsed = round(time.time() - start_time, 3)
            logger.error("Chat error after %.3fs: %s", elapsed, str(e), exc_info=True)
            return ChatResponse(
                text=f"Xin lỗi, đã xảy ra lỗi khi xử lý câu hỏi: {str(e)}",
                display="text",
                metadata={
                    "error": True,
                    "error_message": str(e),
                    "execution_time_seconds": elapsed,
                },
            )

    @staticmethod
    async def export_excel(title: str, sql: Optional[str], data: Optional[List[Dict[str, Any]]], db: Any) -> Any:
        import io
        import re
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        rows = []
        if data and len(data) > 0:
            rows = data
        elif sql:
            rows = await execute_query(sql, {}, db)

        wb = openpyxl.Workbook()
        ws = wb.active
        clean_sheet_title = (title or "HrAi_Export")[:30]
        ws.title = re.sub(r'[^\w\s]', '', clean_sheet_title) or "HrAi_Export"

        if not rows:
            ws.append(["Thông báo", "Không tìm thấy dữ liệu nào phù hợp để xuất Excel."])
        else:
            headers = list(rows[0].keys())

            # Title Header
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
            title_cell = ws.cell(row=1, column=1, value=f"BÁO CÁO DỮ LIỆU HR AI: {(title or 'XUẤT BẢNG DỮ LIỆU').upper()}")
            title_cell.font = Font(name="Arial", size=13, bold=True, color="C00000")
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 32

            # Column Headers
            header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="D9534F", end_color="D9534F", fill_type="solid")
            header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

            ws.row_dimensions[3].height = 26
            for col_idx, h_name in enumerate(headers, 1):
                cell = ws.cell(row=3, column=col_idx, value=str(h_name).replace("_", " ").title())
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align

            # Data Rows
            data_font = Font(name="Arial", size=10)
            thin_border = Border(
                left=Side(style='thin', color='E0E0E0'),
                right=Side(style='thin', color='E0E0E0'),
                top=Side(style='thin', color='E0E0E0'),
                bottom=Side(style='thin', color='E0E0E0')
            )

            for r_idx, row_dict in enumerate(rows, 4):
                ws.row_dimensions[r_idx].height = 20
                for c_idx, h_name in enumerate(headers, 1):
                    val = row_dict.get(h_name)
                    cell = ws.cell(row=r_idx, column=c_idx, value="" if val is None else str(val))
                    cell.font = data_font
                    cell.border = thin_border
                    if r_idx % 2 == 1:
                        cell.fill = PatternFill(start_color="F9F9F9", end_color="F9F9F9", fill_type="solid")

            # Column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    async def run_sql(request: RunSQLRequest, db: Any) -> ChatResponse:
        start_time = time.time()
        sql = request.sql.strip()

        validation_error = validate_sql(sql)
        if validation_error:
            return ChatResponse(
                text=f"Lỗi bảo mật: {validation_error}",
                display="text",
                sql=sql,
                metadata={"error": True, "error_message": validation_error},
            )

        try:
            sql_result = await execute_query(sql, {}, db)
            elapsed = round(time.time() - start_time, 3)
            
            display = "table"
            if len(sql_result) == 1:
                display = "card"
            elif not sql_result:
                display = "text"

            return ChatResponse(
                text="Thực thi SQL thành công." if sql_result else "Không tìm thấy dữ liệu phù hợp.",
                display=display,
                data=sql_result if sql_result else None,
                sql=sql,
                metadata={
                    "row_count": len(sql_result),
                    "execution_time_seconds": elapsed,
                },
            )
        except Exception as e:
            elapsed = round(time.time() - start_time, 3)
            logger.error("Run SQL error: %s", str(e))
            return ChatResponse(
                text=f"Lỗi truy vấn dữ liệu: {str(e)}",
                display="text",
                sql=sql,
                metadata={
                    "error": True,
                    "error_message": str(e),
                    "execution_time_seconds": elapsed,
                },
            )

    @staticmethod
    async def get_sheets(db: Any) -> List[Dict[str, Any]]:
        meta_table_sql = """
        CREATE TABLE IF NOT EXISTS "_sheet_metadata" (
            table_name TEXT PRIMARY KEY,
            sheet_url TEXT,
            tab_name TEXT,
            excluded_columns TEXT,
            accept_inferred_field INTEGER,
            inferred_proposal TEXT,
            last_synced TEXT
        );
        """
        try:
            await db_execute(db, text(meta_table_sql))
            await db_execute(db, text('DELETE FROM "_sheet_metadata" WHERE LOWER(table_name) IN ("users", "accounts", "schedules", "positions", "overtime_requests", "schedule_period");'))
            await db_commit(db)
        except Exception:
            pass

        try:
            meta_res = await db_execute(
                db,
                text('SELECT table_name, sheet_url, tab_name, last_synced FROM "_sheet_metadata"')
            )
            rows = [r for r in meta_res.fetchall() if r[0].lower() not in SYSTEM_TABLE_NAMES]
        except Exception as e:
            logger.error(f"Lỗi lấy danh sách link Google Sheet từ metadata: {e}")
            rows = []

        tabs = []
        tab_id_counter = 1
        last_synced_global = None
        is_sqlite = "sqlite" in str(db.bind.url) if hasattr(db, "bind") and db.bind else True

        for row in rows:
            t_name, sheet_url, tab_name_val, tab_last_synced = row[0], row[1], row[2], row[3]
            row_count = 0
            try:
                cnt_res = await db_execute(db, text(f'SELECT COUNT(*) FROM "{t_name}"'))
                row_count = cnt_res.scalar() or 0
            except Exception:
                continue

            columns = []
            try:
                if is_sqlite:
                    cols_res = await db_execute(db, text(f'PRAGMA table_info("{t_name}")'))
                    for col in cols_res.fetchall():
                        c_name = col[1]
                        c_type = col[2] or "TEXT"
                        if c_name.lower() != "id":
                            columns.append({"name": c_name, "type": c_type.lower()})
                else:
                    cols_res = await db_execute(db, text(
                        "SELECT column_name, data_type FROM information_schema.columns WHERE table_name = :t AND column_name != 'id';"
                    ), {"t": t_name})
                    for col in cols_res.fetchall():
                        columns.append({"name": col[0], "type": col[1].lower()})
            except Exception:
                pass

            clean_title = tab_name_val or t_name.replace("sheet_", "").replace("_", " ").title()
            if tab_last_synced and (not last_synced_global or tab_last_synced > last_synced_global):
                last_synced_global = tab_last_synced

            tabs.append({
                "id": tab_id_counter,
                "sheet_name": f"Google Sheet ({clean_title})",
                "sheet_url": sheet_url or f"https://docs.google.com/spreadsheets/table/{t_name}",
                "table_name": t_name,
                "tab_name": clean_title,
                "row_count": row_count,
                "sync_status": "synced",
                "columns": columns,
                "last_synced": tab_last_synced
            })
            tab_id_counter += 1

        if not tabs:
            return []

        return [{
            "sheet_url": tabs[0]["sheet_url"] if tabs else "",
            "sheet_name": "Danh Sách Link Google Sheet Đã Kết Nối",
            "database_name": "hr_ai",
            "last_synced": last_synced_global or time.strftime("%Y-%m-%d %H:%M:%S"),
            "tabs": tabs
        }]

    @staticmethod
    async def get_db_tables(db: Any) -> List[Dict[str, Any]]:
        meta_table_sql = """
        CREATE TABLE IF NOT EXISTS "_sheet_metadata" (
            table_name TEXT PRIMARY KEY,
            sheet_url TEXT,
            tab_name TEXT,
            excluded_columns TEXT,
            accept_inferred_field INTEGER,
            inferred_proposal TEXT,
            last_synced TEXT
        );
        """
        try:
            await db_execute(db, text(meta_table_sql))
        except Exception:
            pass

        try:
            res = await db_execute(db, text('SELECT table_name, sheet_url, tab_name FROM "_sheet_metadata"'))
            rows = [r for r in res.fetchall() if r[0].lower() not in SYSTEM_TABLE_NAMES]
            meta_map = {row[0]: {"sheet_url": row[1], "tab_name": row[2]} for row in rows}
            tables = list(meta_map.keys())
        except Exception as e:
            logger.error(f"Lỗi lấy danh sách bảng CSDL AI HR: {e}")
            tables = []

        is_sqlite = "sqlite" in str(db.bind.url) if hasattr(db, "bind") and db.bind else True
        result = []
        for t in tables:
            row_count = 0
            try:
                cnt_res = await db_execute(db, text(f'SELECT COUNT(*) FROM "{t}"'))
                row_count = cnt_res.scalar() or 0
            except Exception:
                continue

            columns = []
            try:
                if is_sqlite:
                    cols_res = await db_execute(db, text(f'PRAGMA table_info("{t}")'))
                    for col in cols_res.fetchall():
                        columns.append({"name": col[1], "type": (col[2] or "TEXT").lower()})
                else:
                    cols_res = await db_execute(db, text(
                        "SELECT column_name, data_type FROM information_schema.columns WHERE table_name = :t;"
                    ), {"t": t})
                    for col in cols_res.fetchall():
                        columns.append({"name": col[0], "type": col[1].lower()})
            except Exception:
                pass

            meta_info = meta_map.get(t, {})
            display_name = meta_info.get("tab_name") or t.replace("sheet_", "").replace("_", " ").title()
            sheet_url = meta_info.get("sheet_url")

            result.append({
                "table_name": t,
                "display_name": display_name,
                "row_count": row_count,
                "column_count": len(columns),
                "columns": columns,
                "is_sheet": True,
                "sheet_url": sheet_url
            })

        return result

    @staticmethod
    async def get_db_table_data(table_name: str, limit: int, db: Any) -> Dict[str, Any]:
        try:
            sql = f'SELECT * FROM "{table_name}" LIMIT {limit}'
            sql_result = await execute_query(sql, {}, db)
            return {"table_name": table_name, "data": sql_result, "total": len(sql_result)}
        except Exception as e:
            logger.error(f"Lỗi lấy dữ liệu bảng '{table_name}': {e}")
            return {"table_name": table_name, "data": [], "total": 0, "error": str(e)}

    @staticmethod
    async def get_sheet_data(tab_id: int, limit: int, db: Any) -> Dict[str, Any]:
        try:
            meta_res = await db_execute(db, text('SELECT table_name, sheet_url, tab_name, last_synced FROM "_sheet_metadata"'))
            rows = [r for r in meta_res.fetchall() if r[0].lower() not in SYSTEM_TABLE_NAMES]
            if 1 <= tab_id <= len(rows):
                table_name, sheet_url, tab_name_val, last_synced = rows[tab_id - 1]
                sql = f'SELECT * FROM "{table_name}" LIMIT {limit}'
                sql_result = await execute_query(sql, {}, db)
                return {
                    "tab_id": tab_id,
                    "table_name": table_name,
                    "sheet_url": sheet_url,
                    "tab_name": tab_name_val or table_name,
                    "last_synced": last_synced,
                    "data": sql_result,
                    "total": len(sql_result)
                }
        except Exception as e:
            logger.error(f"Lỗi lấy dữ liệu sheet: {e}")
    @staticmethod
    async def upload_sheet(request_data: Any, db: Any) -> Dict[str, Any]:
        """Tải và nhập dữ liệu Google Sheet trực tiếp từ đường dẫn Link URL hoặc JSON Payload."""
        if isinstance(request_data, str):
            request_data = {
                "sheet_url": request_data.strip(),
                "tab_name": "",
                "excluded_columns": [],
                "accept_inferred_field": False,
                "inferred_proposal": {}
            }
        
        sheet_url = request_data.get("sheet_url", "")
        if not sheet_url or not str(sheet_url).strip():
            raise HTTPException(status_code=400, detail="Vui lòng nhập đường dẫn Link Google Sheet!")

        try:
            result = await create_table_and_sync_sheet(request_data, db)
            return {
                "message": f"🎉 Kết nối và nạp thành công {result.get('row_count', 0)} dòng từ Google Sheet vào bảng '{result.get('table_name', 'sheet_table')}'!",
                "details": result
            }
        except Exception as e:
            logger.error(f"Lỗi kết nối Google Sheet: {e}")
            raise HTTPException(
                status_code=400,
                detail=f"Không thể kết nối Google Sheet. Vui lòng kiểm tra lại đường dẫn và đảm bảo đã bật chế độ chia sẻ 'Bất kỳ ai có liên kết đều có thể xem'. Lỗi: {str(e)}"
            )

    @staticmethod
    async def sync_sheet(tab_id: int, req_data_body: Dict[str, Any], db: Any) -> Dict[str, Any]:
        meta_table_sql = """
        CREATE TABLE IF NOT EXISTS "_sheet_metadata" (
            table_name TEXT PRIMARY KEY,
            sheet_url TEXT,
            tab_name TEXT,
            excluded_columns TEXT,
            accept_inferred_field INTEGER,
            inferred_proposal TEXT,
            last_synced TEXT
        );
        """
        await db_execute(db, text(meta_table_sql))
        await db_commit(db)

        meta_res = await db_execute(
            db,
            text('SELECT table_name, sheet_url, tab_name, excluded_columns, accept_inferred_field, inferred_proposal FROM "_sheet_metadata"')
        )
        rows = meta_res.fetchall()

        if not (1 <= tab_id <= len(rows)):
            raise HTTPException(status_code=404, detail="Không tìm thấy bảng CSDL Sheet/Excel cần đồng bộ.")

        table_name, sheet_url, tab_name_val, excluded_cols_str, accept_inferred_int, inferred_prop_str = rows[tab_id - 1]

        req_data_body = req_data_body or {}
        sheet_url = req_data_body.get("sheet_url") or sheet_url

        if not sheet_url:
            raise HTTPException(
                status_code=400, 
                detail=f"Bảng '{table_name}' chưa có thông tin link Google Sheet. Vui lòng dán lại link Google Sheet và bấm Kết nối."
            )

        if sheet_url.startswith("excel://"):
            import datetime
            now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            await db_execute(
                db, 
                text('UPDATE "_sheet_metadata" SET last_synced = :now WHERE table_name = :tn'),
                {"now": now_str, "tn": table_name}
            )
            await db_commit(db)
            return {
                "message": f"Bảng '{table_name}' là dữ liệu từ file Excel local. Để cập nhật dữ liệu mới từ Excel, vui lòng sử dụng nút 'Upload File Excel'.",
                "table_name": table_name
            }

        excluded_columns = json.loads(excluded_cols_str) if excluded_cols_str else []
        accept_inferred = bool(accept_inferred_int)
        inferred_proposal = json.loads(inferred_prop_str) if inferred_prop_str else {}

        request_data = {
            "sheet_url": sheet_url,
            "tab_name": tab_name_val or table_name.replace("sheet_", ""),
            "excluded_columns": excluded_columns,
            "accept_inferred_field": accept_inferred,
            "inferred_proposal": inferred_proposal,
        }

        try:
            result = await create_table_and_sync_sheet(request_data, db)
            return {
                "message": f"🎉 Đã đồng bộ thành công {result.get('row_count', 0)} dòng mới nhất vào bảng '{result.get('table_name', table_name)}'!",
                "details": result
            }
        except Exception as e:
            logger.error("Lỗi khi đồng bộ Google Sheet: %s", e)
            raise HTTPException(
                status_code=400,
                detail=f"Không thể kết nối Google Sheet. Vui lòng kiểm tra lại đường dẫn Sheet và quyền truy cập (Đặt Chế độ 'Bất kỳ ai có liên kết đều xem được'). Lỗi: {str(e)}"
            )

    @staticmethod
    async def preview_excel(file_bytes: bytes, filename: str) -> Dict[str, Any]:
        try:
            from services.hrai.sheet_pipeline import preprocess_excel_raw_data
            raw_info = preprocess_excel_raw_data(file_bytes, filename)
            analyzed_columns = await analyze_columns_with_gemini_stage1(raw_info["columns"])
            section_analysis = await analyze_sections_with_gemini_stage2(raw_info["sections"], raw_info["sheet_title"], raw_info["tab_name"])
            dynamic_ranges = calculate_dynamic_section_ranges(raw_info["sections"], raw_info["total_rows"])

            headers = [
                {
                    "name": col["field_name"],
                    "data_type": col["data_type"],
                    "description": col["description"],
                    "semantic_features": col.get("semantic_features", {}),
                    "is_inferred": False
                }
                for col in analyzed_columns
            ]

            sample_rows = []
            for r_idx in range(min(5, len(analyzed_columns[0]["sample_values"]) if analyzed_columns else 0)):
                row_sample = [col["sample_values"][r_idx] if r_idx < len(col["sample_values"]) else "" for col in analyzed_columns]
                sample_rows.append(row_sample)

            proposal = section_analysis.get("proposal") or {}
            dividers = [s["text"] for s in raw_info["sections"]]
            if not dividers and proposal.get("mapping"):
                dividers = list(proposal.get("mapping", {}).keys())

            tabs = [{
                "tab_name": raw_info["tab_name"],
                "row_count": raw_info["total_rows"],
                "header_level": 2 if raw_info["is_multi_header"] else 1,
                "section_dividers": dividers,
                "headers": headers,
                "sample_rows": sample_rows,
                "dynamic_ranges": dynamic_ranges,
                "inferred_proposal": proposal if proposal else None
            }]

            return {
                "sheet_url": f"excel://{filename}",
                "sheet_name": f"File Excel ({filename})",
                "tabs": tabs
            }
        except Exception as e:
            logger.error(f"Lỗi preview file Excel: {e}")
            raise HTTPException(
                status_code=400, 
                detail=f"Không thể đọc file Excel '{filename}': {str(e)}"
            )

    @staticmethod
    async def upload_excel(filename: str, file_bytes: bytes, request_data: Dict[str, Any], db: Any) -> Dict[str, Any]:
        try:
            from services.hrai.sheet_pipeline import preprocess_excel_raw_data, is_summary_row
            from services.hrai.sheet_executor import generate_clean_table_name, sanitize_column_name
            raw_info = preprocess_excel_raw_data(file_bytes, filename)
            all_values = raw_info.get("all_values", [])
            if not all_values:
                raise ValueError("File Excel không có dữ liệu!")

            clean_table_name = generate_clean_table_name(raw_info.get("tab_name") or raw_info.get("sheet_title"))

            raw_columns = raw_info.get("columns", [])
            start_data_idx = raw_info.get("start_data_idx", 1)
            data_rows = all_values[start_data_idx:]
            excluded_columns = request_data.get("excluded_columns", [])
            accept_inferred = bool(request_data.get("accept_inferred_field", False))

            headers = []
            seen_sql_names = set()
            for col in raw_columns:
                c_idx = col["column_index"]
                orig_name = col["field_name"]
                if str(c_idx) in excluded_columns or orig_name in excluded_columns:
                    continue

                base_sql_name = sanitize_column_name(orig_name)
                sql_name = base_sql_name
                counter = 1
                while sql_name.lower() in seen_sql_names:
                    sql_name = f"{base_sql_name}_{counter}"
                    counter += 1
                seen_sql_names.add(sql_name.lower())

                col_type = "TEXT"
                if "stt" in orig_name.lower() or orig_name.lower() == "id":
                    col_type = "INTEGER"

                headers.append({"idx": c_idx, "original_name": orig_name, "sql_name": sql_name, "type": col_type})

            is_sqlite = "sqlite" in str(db.bind.url) if hasattr(db, "bind") and db.bind else True
            id_col_sql = "id INTEGER PRIMARY KEY AUTOINCREMENT" if is_sqlite else "id SERIAL PRIMARY KEY"

            columns_sql = [id_col_sql]
            for h in headers:
                columns_sql.append(f'"{h["sql_name"]}" {h["type"]}')

            inferred_proposal = request_data.get("inferred_proposal") or {}
            raw_inferred_name = inferred_proposal.get("field_name", "loai_nhan_su")
            inferred_field_name = sanitize_column_name(raw_inferred_name)
            default_val = inferred_proposal.get("default_value") or inferred_proposal.get("mapping", {}).get("__default__", "TTS")
            mapping = inferred_proposal.get("mapping", {})

            if accept_inferred and inferred_field_name:
                columns_sql.append(f'"{inferred_field_name}" TEXT')

            drop_sql = f'DROP TABLE IF EXISTS "{clean_table_name}"'
            await db_execute(db, text(drop_sql))
            create_sql = f'CREATE TABLE "{clean_table_name}" ({", ".join(columns_sql)})'
            await db_execute(db, text(create_sql))

            sections = raw_info.get("sections", [])
            section_titles = [s["text"] for s in sections]
            curr_status = default_val

            inserted_count = 0
            for r in data_rows:
                non_empty = [str(c).strip() for c in r if str(c).strip()]
                if not non_empty:
                    continue

                if is_summary_row(r):
                    continue

                if len(non_empty) == 1 and non_empty[0] in section_titles:
                    sec_title = non_empty[0]
                    curr_status = mapping.get(sec_title, sec_title)
                    continue

                row_dict = {}
                for h in headers:
                    val = r[h["idx"]] if h["idx"] < len(r) else ""
                    if h["type"] == "INTEGER":
                        try:
                            val = int(val) if val else None
                        except ValueError:
                            val = None
                    row_dict[h["sql_name"]] = val

                if accept_inferred and inferred_field_name:
                    row_dict[inferred_field_name] = curr_status

                col_clause = ", ".join([f'"{k}"' for k in row_dict.keys()])
                param_clause = ", ".join([f":{k}" for k in row_dict.keys()])
                insert_sql = f'INSERT INTO "{clean_table_name}" ({col_clause}) VALUES ({param_clause})'
                await db_execute(db, text(insert_sql), row_dict)
                inserted_count += 1

            meta_table_sql = """
            CREATE TABLE IF NOT EXISTS "_sheet_metadata" (
                table_name TEXT PRIMARY KEY,
                sheet_url TEXT,
                tab_name TEXT,
                excluded_columns TEXT,
                accept_inferred_field INTEGER,
                inferred_proposal TEXT,
                last_synced TEXT
            );
            """
            await db_execute(db, text(meta_table_sql))
            now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            upsert_sql = """
            INSERT INTO "_sheet_metadata" (table_name, sheet_url, tab_name, excluded_columns, accept_inferred_field, inferred_proposal, last_synced)
            VALUES (:table_name, :sheet_url, :tab_name, :excluded_columns, :accept_inferred_field, :inferred_proposal, :last_synced)
            ON CONFLICT(table_name) DO UPDATE SET
                sheet_url=excluded.sheet_url,
                tab_name=excluded.tab_name,
                excluded_columns=excluded.excluded_columns,
                accept_inferred_field=excluded.accept_inferred_field,
                inferred_proposal=excluded.inferred_proposal,
                last_synced=excluded.last_synced;
            """
            await db_execute(db, text(upsert_sql), {
                "table_name": clean_table_name,
                "sheet_url": f"excel://{filename}",
                "tab_name": raw_info["tab_name"],
                "excluded_columns": json.dumps(excluded_columns),
                "accept_inferred_field": 1 if accept_inferred else 0,
                "inferred_proposal": json.dumps(inferred_proposal),
                "last_synced": now_str
            })

            await db_commit(db)
            return {
                "message": f"Đã nhập dữ liệu thành công từ file Excel '{filename}' ({inserted_count} dòng) vào bảng '{clean_table_name}'!",
                "details": {"table_name": clean_table_name, "row_count": inserted_count}
            }
        except Exception as e:
            logger.error(f"Lỗi nhập file Excel: {e}")
            raise HTTPException(status_code=400, detail=f"Không thể xử lý file Excel '{filename}': {str(e)}")

    @staticmethod
    async def delete_sheet(tab_id: int, db: Any) -> Dict[str, Any]:
        try:
            res = await db_execute(db, text('SELECT table_name FROM "_sheet_metadata"'))
            rows = res.fetchall()
            if 1 <= tab_id <= len(rows):
                table_name = rows[tab_id - 1][0]
                await db_execute(db, text(f'DROP TABLE IF EXISTS "{table_name}"'))
                await db_execute(db, text('DELETE FROM "_sheet_metadata" WHERE table_name = :t'), {"t": table_name})
                await db_commit(db)
                return {"message": f"Đã xóa bảng '{table_name}' thành công."}
        except Exception as e:
            logger.error(f"Lỗi xóa bảng CSDL: {e}")
        return {"message": "Đã xóa kết nối thành công."}

    @staticmethod
    def get_settings_data() -> Dict[str, Any]:
        get_settings.cache_clear()
        settings = get_settings()
        api_key = settings.GEMINI_API_KEY or os.environ.get("GEMINI_API_KEY", "")
        masked_key = (api_key[:6] + "..." + api_key[-4:]) if len(api_key) > 10 else api_key
        return {
            "api_key": api_key,
            "masked_api_key": masked_key,
            "model": settings.GEMINI_MODEL or os.environ.get("GEMINI_MODEL", "gemini-3.1-flash-lite")
        }

    @staticmethod
    def save_settings_data(api_key: str, model: str) -> Dict[str, Any]:
        try:
            env_dict = {}
            if os.path.exists(ENV_PATH):
                with open(ENV_PATH, "r", encoding="utf-8") as f:
                    for line in f:
                        line = line.strip()
                        if line and not line.startswith("#") and "=" in line:
                            k, v = line.split("=", 1)
                            env_dict[k.strip()] = v.strip()

            if api_key:
                env_dict["GEMINI_API_KEY"] = api_key
            if model:
                env_dict["GEMINI_MODEL"] = model

            with open(ENV_PATH, "w", encoding="utf-8") as f:
                for k, v in env_dict.items():
                    f.write(f"{k}={v}\n")

            os.environ["GEMINI_API_KEY"] = api_key
            os.environ["GEMINI_MODEL"] = model

            get_settings.cache_clear()
            try:
                from services.hrai.ai_pipeline import llm_sql_generator
                llm_sql_generator._client = None
            except Exception:
                pass

            return {"message": "Đã lưu cài đặt vào .env thành công!"}
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Không thể lưu file .env: {str(e)}")

    @staticmethod
    def test_single_key_data(api_key: str, model_id: str) -> Dict[str, Any]:
        if not api_key:
            return {"working": False, "error": "API Key không được để trống"}

        start_time = time.time()
        m_id = model_id or "gemini-3.1-flash-lite"
        try:
            res_text = ""
            try:
                from google import genai
                client = genai.Client(api_key=api_key)
                response = client.models.generate_content(
                    model=m_id,
                    contents="Xin chào, vui lòng phản hồi ngắn gọn OK."
                )
                res_text = response.text or "OK"
            except Exception:
                import google.generativeai as legacy_genai
                legacy_genai.configure(api_key=api_key)
                model_obj = legacy_genai.GenerativeModel(m_id)
                response = model_obj.generate_content("Xin chào, vui lòng phản hồi ngắn gọn OK.")
                res_text = response.text or "OK"

            latency = int((time.time() - start_time) * 1000)
            return {
                "working": True,
                "latency_ms": latency,
                "response": res_text[:100],
                "message": f"🎉 Kết nối thành công tới Gemini API ({m_id})! Độ trễ: {latency}ms"
            }
        except Exception as e:
            return {
                "working": False,
                "error": f"Lỗi kết nối Gemini API: {str(e)}"
            }

    @staticmethod
    async def generate_schema_proposal(question: str, db: Any) -> Dict[str, Any]:
        """Phân tích câu hỏi của người dùng và tạo proposal cấu trúc bảng (Scenario 1, 2, hoặc 3)"""
        from services.hrai.query_planner_engine.query_planner import _get_genai_client, _call_gemini_generate

        client = _get_genai_client()
        system_prompt = """Bạn là chuyên gia thiết kế cơ sở dữ liệu và phân tích dữ liệu HR Viettel.
NHIỆM VỤ: Phân tích yêu cầu tạo CSDL/Excel/Google Sheet của người dùng và trả về JSON duy nhất.

XÁC ĐỊNH SCENARIO:
- Scenario 1: Dữ liệu hoàn toàn có sẵn trong CSDL (như danh sách nhân sự, thực tập sinh, lịch làm việc, đăng ký OT).
- Scenario 2: Kết hợp cột CSDL sẵn có (họ tên, mã NV, vị trí, sđt) + các cột tùy chỉnh mới chưa có trong CSDL (như điểm KPI, đánh giá hàng tháng, ghi chú mentor).
- Scenario 3: Cấu trúc hoàn toàn mới không thuộc CSDL sẵn có (như danh sách thiết bị dự án, bảng điểm OKR...).

OUTPUT FORMAT (JSON):
{
  "scenario": 1,
  "table_name": "danh_gia_kpi_tts",
  "title": "Bảng Đánh Giá KPI Thực Tập Sinh",
  "reasoning": "Tạo bảng đánh giá KPI cho thực tập sinh",
  "columns": [
    {
      "field_name": "ma_nv",
      "label": "Mã NV",
      "data_type": "TEXT",
      "default_value": "",
      "source_column": "employee_code"
    },
    {
      "field_name": "ho_va_ten",
      "label": "Họ và Tên",
      "data_type": "TEXT",
      "default_value": "",
      "source_column": "full_name"
    },
    {
      "field_name": "diem_kpi",
      "label": "Điểm KPI",
      "data_type": "REAL",
      "default_value": "8.5",
      "source_column": null
    }
  ],
  "sample_rows": [
    { "ma_nv": "TTS01", "ho_va_ten": "Nguyễn Văn A", "diem_kpi": 9.0 }
  ]
}
"""
        prompt = f"YÊU CẦU CỦA NGƯỜI DÙNG: \"{question}\""

        try:
            raw_text = _call_gemini_generate(client, system_prompt, prompt, is_json=True)
            proposal = json.loads(raw_text)

            scenario = proposal.get("scenario", 3)
            cols = proposal.get("columns", [])

            if scenario in (1, 2):
                try:
                    res = await db_execute(db, text("SELECT employee_code, full_name, role, project, phone FROM users LIMIT 5;"))
                    db_rows = res.fetchall()
                    if db_rows:
                        sample_rows = []
                        for db_r in db_rows:
                            row_dict = {}
                            db_map = {
                                "ma_nv": db_r[0], "employee_code": db_r[0],
                                "ho_va_ten": db_r[1], "full_name": db_r[1],
                                "vi_tri": db_r[2], "role": db_r[2], "position": db_r[2],
                                "du_an": db_r[3], "project": db_r[3],
                                "so_dien_thoai": db_r[4], "phone": db_r[4]
                            }
                            for c in cols:
                                fname = c["field_name"]
                                if fname in db_map:
                                    row_dict[fname] = db_map[fname]
                                elif c.get("source_column") and c["source_column"] in db_map:
                                    row_dict[fname] = db_map[c["source_column"]]
                                else:
                                    row_dict[fname] = c.get("default_value") or ""
                            sample_rows.append(row_dict)
                        if sample_rows:
                            proposal["sample_rows"] = sample_rows
                except Exception as e:
                    logger.warning("Failed to enrich proposal sample rows: %s", e)

            return proposal
        except Exception as e:
            logger.error("Failed to generate schema proposal: %s", e)
            return {
                "scenario": 3,
                "table_name": "custom_table_" + str(int(time.time())),
                "title": "Bảng Dữ Liệu Tùy Chỉnh",
                "reasoning": f"Tự động khởi tạo từ câu hỏi: {question}",
                "columns": [
                    {"field_name": "stt", "label": "STT", "data_type": "INTEGER", "default_value": "1"},
                    {"field_name": "ho_ten", "label": "Họ và Tên", "data_type": "TEXT", "default_value": ""},
                    {"field_name": "noi_dung", "label": "Nội Dung", "data_type": "TEXT", "default_value": ""}
                ],
                "sample_rows": [
                    {"stt": 1, "ho_ten": "Nguyễn Văn A", "noi_dung": "Dữ liệu mẫu 1"}
                ]
            }

    @staticmethod
    async def create_custom_table(table_name: str, title: str, columns: List[Dict[str, Any]], rows: List[Dict[str, Any]], db: Any) -> Dict[str, Any]:
        """Tạo bảng SQLite thực tế, nạp các dòng dữ liệu và lưu metadata vào _sheet_metadata"""
        import re
        clean_table_name = re.sub(r'[^a-zA-Z0-9_]', '', table_name.lower().strip())
        if not clean_table_name.startswith("table_") and not clean_table_name.startswith("sheet_"):
            clean_table_name = "sheet_" + clean_table_name

        if not columns:
            raise HTTPException(status_code=400, detail="Danh sách cột không được để trống!")

        col_defs = ["id INTEGER PRIMARY KEY AUTOINCREMENT"]
        clean_cols = []
        for c in columns:
            fname = re.sub(r'[^a-zA-Z0-9_]', '', str(c.get("field_name", "")).lower().strip())
            if not fname or fname == "id":
                continue
            dtype = str(c.get("data_type", "TEXT")).upper()
            if dtype not in ("TEXT", "INTEGER", "REAL", "DATE"):
                dtype = "TEXT"
            col_defs.append(f'"{fname}" {dtype}')
            clean_cols.append({"name": fname, "label": c.get("label", fname), "type": dtype})

        create_sql = f'CREATE TABLE IF NOT EXISTS "{clean_table_name}" (\n  ' + ',\n  '.join(col_defs) + '\n);'
        await db_execute(db, text(create_sql))

        insert_count = 0
        if rows:
            col_names = [c["name"] for c in clean_cols]
            param_placeholders = ", ".join([f":{col}" for col in col_names])
            insert_sql = f'INSERT INTO "{clean_table_name}" (' + ', '.join([f'"{col}"' for col in col_names]) + f') VALUES ({param_placeholders})'

            for r in rows:
                row_params = {}
                for col in col_names:
                    row_params[col] = r.get(col) if r.get(col) is not None else ""
                try:
                    await db_execute(db, text(insert_sql), row_params)
                    insert_count += 1
                except Exception as e:
                    logger.warning("Row insert warning in %s: %s", clean_table_name, e)

        meta_table_sql = """
        CREATE TABLE IF NOT EXISTS "_sheet_metadata" (
            table_name TEXT PRIMARY KEY,
            sheet_url TEXT,
            tab_name TEXT,
            excluded_columns TEXT,
            accept_inferred_field INTEGER,
            inferred_proposal TEXT,
            last_synced TEXT
        );
        """
        await db_execute(db, text(meta_table_sql))

        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        meta_insert_sql = """
        INSERT INTO "_sheet_metadata" (table_name, sheet_url, tab_name, last_synced)
        VALUES (:t_name, :s_url, :t_title, :l_synced)
        ON CONFLICT(table_name) DO UPDATE SET
            tab_name=EXCLUDED.tab_name,
            last_synced=EXCLUDED.last_synced;
        """
        await db_execute(db, text(meta_insert_sql), {
            "t_name": clean_table_name,
            "s_url": "",
            "t_title": title or clean_table_name,
            "l_synced": now_str
        })
        await db_commit(db)

        return {
            "success": True,
            "table_name": clean_table_name,
            "title": title,
            "row_count": insert_count,
            "column_count": len(clean_cols),
            "message": f"🎉 Đã lưu thành công bảng '{title}' vào CSDL AI HR ({insert_count} bản ghi)!"
        }

    @staticmethod
    async def auto_create_google_sheet_for_table(table_name: str, db: Any) -> Dict[str, Any]:
        """Tạo Google Sheet tự động cho một bảng CSDL trong AI HR CSDL"""
        from services.google_sheets_service import GoogleSheetsService

        meta_res = await db_execute(db, text('SELECT sheet_url, tab_name FROM "_sheet_metadata" WHERE table_name = :t'), {"t": table_name})
        meta_row = meta_res.fetchone()
        existing_sheet_url = meta_row[0] if meta_row else None
        title = meta_row[1] if meta_row and meta_row[1] else table_name

        cols_res = await db_execute(db, text(f'PRAGMA table_info("{table_name}")'))
        headers = [col[1] for col in cols_res.fetchall() if col[1].lower() != "id"]

        data_res = await db_execute(db, text(f'SELECT * FROM "{table_name}" LIMIT 500'))
        raw_rows = [dict(r._mapping) if hasattr(r, "_mapping") else dict(r) for r in data_res.fetchall()]

        sheet_result = GoogleSheetsService.auto_create_generic_table_sheet(
            title=title,
            headers=headers,
            rows_data=raw_rows,
            target_sheet_url=existing_sheet_url
        )

        if sheet_result.get("success") and sheet_result.get("sheet_url"):
            now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            await db_execute(db, text('UPDATE "_sheet_metadata" SET sheet_url = :url, last_synced = :sync WHERE table_name = :t'), {
                "url": sheet_result["sheet_url"],
                "sync": now_str,
                "t": table_name
            })
            await db_commit(db)

        return sheet_result
