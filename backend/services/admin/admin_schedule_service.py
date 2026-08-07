import calendar
from datetime import datetime, date
from io import BytesIO
from typing import List, Optional
from fastapi import HTTPException
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session

import models
import schemas


def natural_sort_key(u):
    code = getattr(u, 'employee_code', '') or ''
    import re
    match = re.search(r'\d+', str(code))
    if match:
        prefix = str(code)[:match.start()].lower()
        num = int(match.group(0))
        return (prefix, num, str(code).lower())
    return (str(code).lower(), 0, str(code).lower())


class AdminScheduleService:
    @staticmethod
    def get_period(month: int, year: int, db: Session) -> models.SchedulePeriod:
        period = db.query(models.SchedulePeriod).filter(
            models.SchedulePeriod.month == month,
            models.SchedulePeriod.year == year
        ).first()
        if not period:
            period = models.SchedulePeriod(month=month, year=year, status="open")
            db.add(period)
            db.commit()
            db.refresh(period)
        return period

    @staticmethod
    def open_period(body: schemas.SchedulePeriodOpen, db: Session) -> models.SchedulePeriod:
        period = AdminScheduleService.get_period(body.month, body.year, db)
        period.status = "open"
        period.open_date = datetime.now()
        period.close_date = body.close_date
        db.commit()
        db.refresh(period)
        return period

    @staticmethod
    def close_period(month: int, year: int, db: Session) -> models.SchedulePeriod:
        period = AdminScheduleService.get_period(month, year, db)
        period.status = "closed"
        db.commit()
        db.refresh(period)
        return period

    @staticmethod
    def get_schedule(month: int, year: int, db: Session) -> dict:
        num_days = calendar.monthrange(year, month)[1]
        period = db.query(models.SchedulePeriod).filter(
            models.SchedulePeriod.month == month,
            models.SchedulePeriod.year == year
        ).first()

        interns = db.query(models.User).filter(
            models.User.user_type == "intern"
        ).all()
        interns.sort(key=natural_sort_key)

        first_day = date(year, month, 1)
        last_day = date(year, month, num_days)

        schedules = db.query(models.Schedule).filter(
            models.Schedule.work_day >= first_day,
            models.Schedule.work_day <= last_day
        ).all()

        schedule_map = {}
        for s in schedules:
            uid = s.user_id
            day = s.work_day.day
            if uid not in schedule_map:
                schedule_map[uid] = {}
            schedule_map[uid][day] = s.shift

        rows = []
        for i in interns:
            user_shifts = schedule_map.get(i.id, {})
            days_dict = {}
            for d in range(1, num_days + 1):
                days_dict[d] = user_shifts.get(d, None)

            rows.append({
                "user_id": i.id,
                "employee_code": i.employee_code,
                "full_name": i.full_name,
                "role": i.role or i.position or "TTS",
                "project": i.project or "—",
                "days": days_dict
            })

        return {
            "month": month,
            "year": year,
            "num_days": num_days,
            "period": {
                "id": period.id if period else None,
                "status": period.status if period else "open",
                "open_date": period.open_date if period else None,
                "close_date": period.close_date if period else None
            },
            "rows": rows
        }

    @staticmethod
    def export_schedule(month: int, year: int, db: Session) -> StreamingResponse:
        data = AdminScheduleService.get_schedule(month, year, db)
        num_days = data["num_days"]
        rows = data["rows"]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Lich_TTS_Thang_{month}_{year}"

        title = f"BẢNG LỊCH LÀM VIỆC THỰC TẬP SINH THÁNG {month}/{year}"
        total_cols = 4 + num_days
        last_col_letter = openpyxl.utils.get_column_letter(total_cols)

        ws.merge_cells(f"A1:{last_col_letter}1")
        c1 = ws["A1"]
        c1.value = title
        c1.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        c1.fill = PatternFill("solid", fgColor="1E3A5F")
        c1.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        headers = ["STT", "MÃ NV", "HỌ VÀ TÊN", "DỰ ÁN"] + [str(d) for d in range(1, num_days + 1)]
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="2563EB")
        center = Alignment(horizontal="center", vertical="center")

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(2, col_idx, h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
        ws.row_dimensions[2].height = 26

        thin = Side(style="thin", color="CBD5E1")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for row_idx, r in enumerate(rows, start=3):
            ws.cell(row_idx, 1, row_idx - 2).alignment = center
            ws.cell(row_idx, 2, r["employee_code"]).alignment = center
            ws.cell(row_idx, 3, r["full_name"]).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row_idx, 4, r["project"]).alignment = Alignment(horizontal="left", vertical="center")

            for d in range(1, num_days + 1):
                shift = r["days"].get(d) or ""
                c = ws.cell(row_idx, 4 + d, shift)
                c.alignment = center
                if shift == "SC":
                    c.fill = PatternFill("solid", fgColor="DCFCE7")
                elif shift in ("S", "C"):
                    c.fill = PatternFill("solid", fgColor="EFF6FF")

            for col in range(1, total_cols + 1):
                ws.cell(row_idx, col).border = border
            ws.row_dimensions[row_idx].height = 20

        for col in ws.columns:
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            if col[0].column in (1, 2):
                ws.column_dimensions[col_letter].width = 12
            elif col[0].column == 3:
                ws.column_dimensions[col_letter].width = 24
            elif col[0].column == 4:
                ws.column_dimensions[col_letter].width = 20
            else:
                ws.column_dimensions[col_letter].width = 6

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"Bang_Lich_Thuc_Tap_Thang_{month}_{year}.xlsx"
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
