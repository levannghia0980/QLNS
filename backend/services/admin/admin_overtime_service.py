import calendar
from datetime import date, datetime
from io import BytesIO
from typing import List, Optional
from fastapi import HTTPException
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session

import models
import schemas


def enrich_ot(ot: models.OvertimeRequest) -> dict:
    d = {c.name: getattr(ot, c.name) for c in ot.__table__.columns}
    if ot.user:
        d["employee_code"] = ot.user.employee_code
        d["full_name"] = ot.user.full_name
    return d


class AdminOvertimeService:
    @staticmethod
    def admin_list_ot(
        month: Optional[int],
        year: Optional[int],
        project: Optional[str],
        status: Optional[str],
        employee_name: Optional[str],
        staff_category: Optional[str],
        db: Session
    ) -> List[dict]:
        now = datetime.now()
        m = month or now.month
        y = year or now.year
        first_day = date(y, m, 1)
        last_day = date(y, m, calendar.monthrange(y, m)[1])

        q = db.query(models.OvertimeRequest).join(
            models.User, models.OvertimeRequest.user_id == models.User.id
        ).filter(
            models.OvertimeRequest.work_date >= first_day,
            models.OvertimeRequest.work_date <= last_day,
        )

        if project:
            q = q.filter(models.OvertimeRequest.project.ilike(f"%{project}%"))
        if status:
            q = q.filter(models.OvertimeRequest.status == status)
        if employee_name:
            q = q.filter(models.User.full_name.ilike(f"%{employee_name}%"))
        if staff_category:
            q = q.filter(models.User.staff_category == staff_category)

        records = q.order_by(models.OvertimeRequest.created_at.desc(), models.OvertimeRequest.work_date.desc(), models.OvertimeRequest.id.desc()).all()
        return [enrich_ot(r) for r in records]

    @staticmethod
    def admin_ot_summary(month: Optional[int], year: Optional[int], project: Optional[str], db: Session) -> dict:
        now = datetime.now()
        m = month or now.month
        y = year or now.year
        first_day = date(y, m, 1)
        last_day = date(y, m, calendar.monthrange(y, m)[1])
        num_days = calendar.monthrange(y, m)[1]

        q = db.query(models.OvertimeRequest).join(
            models.User, models.OvertimeRequest.user_id == models.User.id
        ).filter(
            models.OvertimeRequest.work_date >= first_day,
            models.OvertimeRequest.work_date <= last_day,
        )
        if project:
            q = q.filter(models.OvertimeRequest.project.ilike(f"%{project}%"))

        records = q.order_by(models.User.full_name, models.OvertimeRequest.work_date, models.OvertimeRequest.start_time).all()
        factors_order = [1.5, 2.1, 2.0, 2.7, 3.0, 3.9]
        user_map: dict = {}
        user_order: list = []

        day_of_week_names = {
            0: "Thứ 2", 1: "Thứ 3", 2: "Thứ 4", 3: "Thứ 5", 4: "Thứ 6", 5: "Thứ 7", 6: "Chủ nhật"
        }

        for ot in records:
            uid = ot.user_id
            if uid not in user_map:
                user_map[uid] = {
                    "user_id": uid,
                    "employee_code": ot.user.employee_code if ot.user else "",
                    "full_name": ot.user.full_name if ot.user else "",
                    "project": ot.project or (ot.user.project if ot.user else ""),
                    "days": {},
                    "detail_records": [],
                    "total_by_factor": {f"{f:.1f}": 0.0 for f in factors_order},
                    "total_raw": 0.0,
                    "total_weighted": 0.0,
                    "total_raw_hours": 0.0,
                    "total_weighted_hours": 0.0,
                    "approved_raw_hours": 0.0,
                    "approved_weighted_hours": 0.0,
                    "pending_raw_hours": 0.0,
                    "pending_weighted_hours": 0.0,
                    "rejected_raw_hours": 0.0,
                    "rejected_weighted_hours": 0.0,
                    "total_count": 0,
                    "pending_count": 0,
                    "approved_count": 0,
                    "rejected_count": 0,
                }
                user_order.append(uid)

            dow_idx = ot.work_date.weekday()
            dow_name = day_of_week_names.get(dow_idx, f"Thứ {dow_idx + 2}")

            record_item = {
                "ot_id": ot.id,
                "work_date": ot.work_date.strftime("%Y-%m-%d"),
                "date_display": ot.work_date.strftime("%d/%m/%Y"),
                "day_number": ot.work_date.day,
                "day_of_week": dow_name,
                "start_time": ot.start_time,
                "end_time": ot.end_time,
                "raw_hours": round(float(ot.raw_hours or 0.0), 2),
                "factor": float(ot.factor or 1.5),
                "weighted_hours": round(float(ot.weighted_hours or 0.0), 2),
                "reason": ot.reason or "",
                "status": ot.status,
                "reject_reason": ot.reject_reason or "",
                "is_holiday": float(ot.factor or 0) >= 3.0,
                "is_weekend": dow_idx in (5, 6),
            }

            day = ot.work_date.day
            if day not in user_map[uid]["days"]:
                user_map[uid]["days"][day] = []
            user_map[uid]["days"][day].append(record_item)
            user_map[uid]["detail_records"].append(record_item)

            factor_key = f"{float(ot.factor):.1f}"
            raw = round(float(ot.raw_hours or 0.0), 2)
            weighted = round(float(ot.weighted_hours or 0.0), 2)

            user_map[uid]["total_count"] += 1
            user_map[uid]["total_raw"] = round(user_map[uid]["total_raw"] + raw, 2)
            user_map[uid]["total_weighted"] = round(user_map[uid]["total_weighted"] + weighted, 2)
            user_map[uid]["total_raw_hours"] = user_map[uid]["total_raw"]
            user_map[uid]["total_weighted_hours"] = user_map[uid]["total_weighted"]

            user_map[uid]["total_by_factor"][factor_key] = round(
                user_map[uid]["total_by_factor"].get(factor_key, 0.0) + raw, 2
            )

            if ot.status == "Approved":
                user_map[uid]["approved_count"] += 1
                user_map[uid]["approved_raw_hours"] = round(user_map[uid]["approved_raw_hours"] + raw, 2)
                user_map[uid]["approved_weighted_hours"] = round(user_map[uid]["approved_weighted_hours"] + weighted, 2)
            elif ot.status == "Pending":
                user_map[uid]["pending_count"] += 1
                user_map[uid]["pending_raw_hours"] = round(user_map[uid]["pending_raw_hours"] + raw, 2)
                user_map[uid]["pending_weighted_hours"] = round(user_map[uid]["pending_weighted_hours"] + weighted, 2)
            elif ot.status == "Rejected":
                user_map[uid]["rejected_count"] += 1
                user_map[uid]["rejected_raw_hours"] = round(user_map[uid]["rejected_raw_hours"] + raw, 2)
                user_map[uid]["rejected_weighted_hours"] = round(user_map[uid]["rejected_weighted_hours"] + weighted, 2)

        return {
            "month": m,
            "year": y,
            "num_days": num_days,
            "rows": [user_map[uid] for uid in user_order],
        }

    @staticmethod
    def admin_approve_ot(ot_id: int, admin: models.User, db: Session) -> dict:
        ot = db.query(models.OvertimeRequest).filter(models.OvertimeRequest.id == ot_id).first()
        if not ot:
            raise HTTPException(status_code=404, detail="Không tìm thấy yêu cầu OT.")
        if ot.status != "Pending":
            raise HTTPException(status_code=400, detail="Chỉ duyệt được OT đang ở trạng thái Pending.")
        ot.status = "Approved"
        ot.approved_by = admin.id
        ot.approved_at = datetime.now()
        ot.reject_reason = None
        db.commit()
        db.refresh(ot)
        return enrich_ot(ot)

    @staticmethod
    def admin_reject_ot(ot_id: int, body: schemas.OvertimeApprove, admin: models.User, db: Session) -> dict:
        if not body.reject_reason or not body.reject_reason.strip():
            raise HTTPException(status_code=400, detail="Vui lòng nhập lý do từ chối.")
        ot = db.query(models.OvertimeRequest).filter(models.OvertimeRequest.id == ot_id).first()
        if not ot:
            raise HTTPException(status_code=404, detail="Không tìm thấy yêu cầu OT.")
        if ot.status != "Pending":
            raise HTTPException(status_code=400, detail="Chỉ từ chối được OT đang ở trạng thái Pending.")
        ot.status = "Rejected"
        ot.reject_reason = body.reject_reason.strip()
        ot.approved_by = admin.id
        ot.approved_at = datetime.now()
        db.commit()
        db.refresh(ot)
        return enrich_ot(ot)

    @staticmethod
    def admin_approve_all_pending(month: Optional[int], year: Optional[int], project: Optional[str], admin: models.User, db: Session) -> dict:
        now = datetime.now()
        m = month or now.month
        y = year or now.year
        first_day = date(y, m, 1)
        last_day = date(y, m, calendar.monthrange(y, m)[1])

        q = db.query(models.OvertimeRequest).filter(
            models.OvertimeRequest.work_date >= first_day,
            models.OvertimeRequest.work_date <= last_day,
            models.OvertimeRequest.status == "Pending",
        )
        if project:
            q = q.filter(models.OvertimeRequest.project.ilike(f"%{project}%"))

        count = 0
        for ot in q.all():
            ot.status = "Approved"
            ot.approved_by = admin.id
            ot.approved_at = datetime.now()
            count += 1
        db.commit()
        return {"message": f"Đã duyệt tất cả {count} yêu cầu OT đang chờ."}

    @staticmethod
    def export_excel(month: Optional[int], year: Optional[int], project: Optional[str], db: Session) -> StreamingResponse:
        summary = AdminOvertimeService.admin_ot_summary(month, year, project, db)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Phu_Luc_02_T{summary['month']}_{summary['year']}"

        title = f"BẢNG TỔNG HỢP LÀM THÊM GIỜ (OT) THÁNG {summary['month']}/{summary['year']}"
        ws.merge_cells("A1:G1")
        c1 = ws["A1"]
        c1.value = title
        c1.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        c1.fill = PatternFill("solid", fgColor="EE0033")
        c1.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        headers = ["STT", "MÃ NHÂN VIÊN", "HỌ VÀ TÊN", "DỰ ÁN", "SỐ BUỔI OT", "TỔNG GIỜ THỰC TẾ", "TỔNG GIỜ QUY ĐỔI"]
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1E293B")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(3, col_idx, h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
        ws.row_dimensions[3].height = 28

        thin = Side(style="thin", color="CBD5E1")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for row_idx, r in enumerate(summary["rows"], start=4):
            ws.cell(row_idx, 1, row_idx - 3).alignment = Alignment(horizontal="center")
            ws.cell(row_idx, 2, r["employee_code"]).alignment = Alignment(horizontal="center")
            ws.cell(row_idx, 3, r["full_name"]).alignment = Alignment(horizontal="left")
            ws.cell(row_idx, 4, r["project"] or "—").alignment = Alignment(horizontal="left")
            ws.cell(row_idx, 5, len(r.get("detail_records", []))).alignment = Alignment(horizontal="center")
            ws.cell(row_idx, 6, r["total_raw_hours"]).alignment = Alignment(horizontal="center")
            ws.cell(row_idx, 7, r["total_weighted_hours"]).alignment = Alignment(horizontal="center")

            for c in range(1, 8):
                ws.cell(row_idx, c).border = border
            ws.row_dimensions[row_idx].height = 22

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        filename = f"Phu_Luc_02_Bao_Cao_OT_Thang_{summary['month']}_{summary['year']}.xlsx"
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
