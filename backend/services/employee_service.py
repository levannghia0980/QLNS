from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from typing import List, Optional
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import date, datetime
import unicodedata
import re

import models
import schemas
import auth


def remove_accents(input_str: str) -> str:
    s = input_str.replace('đ', 'd').replace('Đ', 'D')
    nfkd_form = unicodedata.normalize('NFKD', s)
    return "".join([c for c in nfkd_form if not unicodedata.combining(c)])


def enrich_employee(user: models.User) -> dict:
    """Bổ sung thông tin position_name cho employee response"""
    d = {c.name: getattr(user, c.name) for c in user.__table__.columns}
    d["position_name"] = user.position_rel.name if user.position_rel else None
    return d


class EmployeeService:
    @staticmethod
    def list_positions(db: Session) -> List[models.Position]:
        return db.query(models.Position).order_by(models.Position.id).all()

    @staticmethod
    def list_managers(db: Session) -> List[schemas.ManagerResponse]:
        users = (
            db.query(models.User)
            .join(models.Position, models.User.position_id == models.Position.id)
            .filter(models.Position.is_manager == True, models.User.user_type == "employee")
            .order_by(models.User.full_name)
            .all()
        )
        result = []
        for u in users:
            username = None
            if u.viettel_email and "@" in u.viettel_email:
                username = u.viettel_email.split("@")[0]
            elif u.account:
                username = u.account.username

            result.append(schemas.ManagerResponse(
                id=u.id,
                full_name=u.full_name,
                username=username,
                position_name=u.position_rel.name if u.position_rel else None,
            ))
        return result

    @staticmethod
    def list_employees(db: Session) -> List[schemas.EmployeeResponse]:
        users = (
            db.query(models.User)
            .filter(models.User.user_type == "employee")
            .order_by(models.User.created_at.desc())
            .all()
        )
        result = []
        for u in users:
            data = enrich_employee(u)
            result.append(schemas.EmployeeResponse(**data))
        return result

    @staticmethod
    def create_employee(data: schemas.EmployeeCreate, db: Session) -> schemas.EmployeeResponse:
        existing = db.query(models.User).filter(models.User.employee_code == data.employee_code).first()
        if existing:
            raise HTTPException(status_code=400, detail="Mã nhân viên đã tồn tại")

        user_data = data.model_dump()
        if user_data.get("staff_category") != "Cho mượn":
            user_data["borrow_end_date"] = None
            user_data["borrow_project"] = None
            user_data["borrow_pm"] = None
            user_data["borrow_center"] = None

        user_data["user_type"] = "employee"
        user_data["role"] = "user"

        if not data.viettel_email:
            raise HTTPException(status_code=400, detail="Vui lòng nhập Email Viettel (được dùng làm tài khoản đăng nhập).")

        existing_acc = db.query(models.Account).filter(models.Account.username == data.viettel_email).first()
        if existing_acc:
            raise HTTPException(status_code=400, detail="Email Viettel này đã được sử dụng.")

        user = models.User(**user_data)
        db.add(user)
        db.flush()

        account = models.Account(
            user_id=user.id,
            username=data.viettel_email,
            password=auth.hash_password("123456"),
        )
        db.add(account)
        db.commit()
        db.refresh(user)

        resp = enrich_employee(user)
        return schemas.EmployeeResponse(**resp)

    @staticmethod
    def update_employee(employee_id: int, data: schemas.EmployeeUpdate, db: Session) -> schemas.EmployeeResponse:
        user = db.query(models.User).filter(
            models.User.id == employee_id, models.User.user_type == "employee"
        ).first()
        if not user:
            raise HTTPException(status_code=404, detail="Không tìm thấy nhân viên")

        update_data = data.model_dump(exclude_unset=True)
        staff_cat = update_data.get("staff_category", user.staff_category)
        if staff_cat != "Cho mượn":
            update_data["borrow_end_date"] = None
            update_data["borrow_project"] = None
            update_data["borrow_pm"] = None
            update_data["borrow_center"] = None

        for field, val in update_data.items():
            setattr(user, field, val)

        db.commit()
        db.refresh(user)
        resp = enrich_employee(user)
        return schemas.EmployeeResponse(**resp)

    @staticmethod
    def delete_employee(employee_id: int, current_user: models.User, db: Session) -> dict:
        user = db.query(models.User).filter(
            models.User.id == employee_id, models.User.user_type == "employee"
        ).first()
        if not user:
            raise HTTPException(status_code=404, detail="Không tìm thấy nhân viên")
        if user.id == current_user.id:
            raise HTTPException(status_code=400, detail="Không thể xóa tài khoản đang đăng nhập")
        db.delete(user)
        db.commit()
        return {"message": f"Đã xóa nhân viên {user.full_name}"}

    @staticmethod
    def download_template() -> StreamingResponse:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Danh sách Nhân sự"

        headers = [
            "HỌ VÀ TÊN (*)", "MÃ NV (*)", "VỊ TRÍ", "GIỚI TÍNH", "DÂN TỘC",
            "EMAIL VIETTEL", "NGÀY SINH (YYYY-MM-DD)", "QUÊ QUÁN", "SỐ ĐIỆN THOẠI",
            "SỐ CCCD", "SERI MÁY TÍNH", "SỐ TÀI KHOẢN", "NGÂN HÀNG",
            "LOẠI NHÂN SỰ (NS trung tâm/Onsite/Cho mượn)",
        ]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1e3a5f")
        center = Alignment(horizontal="center", vertical="center")

        for col_idx, header in enumerate(headers, start=1):
            c = ws.cell(1, col_idx, header)
            c.font = header_font
            c.fill = header_fill
            c.alignment = center
            ws.column_dimensions[c.column_letter].width = 22

        sample = ["Nguyễn Văn A", "NV001", "Dev", "Nam", "Kinh",
                  "nguyenvana@viettel.com.vn", "1995-01-15", "Hà Nội",
                  "0901234567", "012345678901", "Laptop Latitude 3420: 8GQBFG3",
                  "1234567890", "Viettel Money", "NS trung tâm"]
        for col_idx, val in enumerate(sample, start=1):
            ws.cell(2, col_idx, val)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=template_import_nhansu.xlsx"},
        )

    @staticmethod
    def import_employees(file_bytes: bytes, db: Session) -> schemas.EmployeeImportResult:
        buf = BytesIO(file_bytes)
        wb = openpyxl.load_workbook(buf, data_only=True)
        ws = wb.active

        success_count = 0
        skip_count = 0
        renamed = []

        has_extra_col_17 = False
        for cell in ws[1]:
            if cell.value and "Dự án đang làm" in str(cell.value):
                if cell.column - 1 == 18:
                    has_extra_col_17 = True
                break

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 3 or not row[1] or not row[2]:
                if row[0] and isinstance(row[0], str) and row[1] and (isinstance(row[1], str) or isinstance(row[1], int)):
                    offset = -1
                else:
                    continue
            else:
                offset = 0

            def get_val(idx):
                actual_idx = idx + offset
                if has_extra_col_17 and idx >= 17:
                    actual_idx += 1
                if actual_idx < 0 or actual_idx >= len(row):
                    return None
                val = row[actual_idx]
                return str(val).strip() if val is not None and str(val).strip() != "" else None

            full_name = get_val(1)
            emp_code = get_val(2)
            if not full_name or not emp_code:
                continue

            existing = db.query(models.User).filter(models.User.employee_code == emp_code).first()
            if existing:
                skip_count += 1
                continue

            birthday = None
            raw_bday = row[9 + offset] if (9 + offset) < len(row) else None
            if raw_bday:
                if isinstance(raw_bday, datetime):
                    birthday = raw_bday.date()
                elif isinstance(raw_bday, date):
                    birthday = raw_bday
                elif isinstance(raw_bday, str):
                    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
                        try:
                            birthday = datetime.strptime(raw_bday.strip(), fmt).date()
                            break
                        except ValueError:
                            pass

            borrow_end = None
            raw_bend = row[21 + offset] if (21 + offset) < len(row) else None
            if raw_bend:
                if isinstance(raw_bend, datetime):
                    borrow_end = raw_bend.date()
                elif isinstance(raw_bend, date):
                    borrow_end = raw_bend
                elif isinstance(raw_bend, str):
                    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
                        try:
                            borrow_end = datetime.strptime(raw_bend.strip(), fmt).date()
                            break
                        except ValueError:
                            pass

            position_id = None
            position_name = get_val(3)
            if position_name:
                pos = db.query(models.Position).filter(func.lower(models.Position.name) == position_name.lower()).first()
                if pos:
                    position_id = pos.id

            is_onsite = str(get_val(19) or "").lower() in ("true", "đúng", "1", "yes", "có")
            is_borrowed = str(get_val(20) or "").lower() in ("true", "đúng", "1", "yes", "có")

            staff_cat = "NS trung tâm"
            if is_borrowed:
                staff_cat = "Cho mượn"
            elif is_onsite:
                staff_cat = "Onsite"

            use_mac = "Có" if str(get_val(16) or "").lower() in ("true", "đúng", "1", "yes", "có") else "Không"

            user_data = {
                "full_name": full_name,
                "employee_code": emp_code,
                "user_type": "employee",
                "role": "user",
                "position_id": position_id,
                "project": get_val(4),
                "direct_manager": get_val(5),
                "gender": get_val(6),
                "ethnicity": get_val(7),
                "viettel_email": get_val(8),
                "birthday": birthday,
                "hometown": get_val(10),
                "phone": get_val(11),
                "cccd": get_val(12),
                "computer_serial": get_val(13),
                "bank_account": get_val(14),
                "bank_name": "Viettel Money",
                "employment_status": get_val(15) or "Thử việc",
                "use_company_mac": use_mac,
                "seat_position": get_val(18),
                "staff_category": staff_cat,
                "borrow_end_date": borrow_end if staff_cat == "Cho mượn" else None,
                "borrow_project": get_val(22) if staff_cat == "Cho mượn" else None,
                "borrow_pm": get_val(23) if staff_cat == "Cho mượn" else None,
                "borrow_center": get_val(24) if staff_cat == "Cho mượn" else None,
                "account_status": 1,
            }

            viettel_email = get_val(8)
            if not viettel_email:
                no_accent = remove_accents(full_name).lower()
                clean_name = re.sub(r'[^a-z0-9]', '', no_accent)
                viettel_email = f"{clean_name}@viettel.com.vn"

            username = viettel_email
            existing_acc = db.query(models.Account).filter(models.Account.username == username).first()
            counter = 1
            original_username = username
            while existing_acc:
                username = f"{original_username.split('@')[0]}{counter}@{original_username.split('@')[1]}"
                existing_acc = db.query(models.Account).filter(models.Account.username == username).first()
                counter += 1

            if username != viettel_email:
                renamed.append(f"{full_name} ({emp_code}): {viettel_email} -> {username}")

            user_data["viettel_email"] = username
            user = models.User(**user_data)
            db.add(user)
            db.flush()

            account = models.Account(
                user_id=user.id,
                username=username,
                password=auth.hash_password("123456"),
            )
            db.add(account)
            success_count += 1

        db.commit()
        return schemas.EmployeeImportResult(
            imported=success_count,
            skipped=skip_count,
            renamed=renamed
        )
